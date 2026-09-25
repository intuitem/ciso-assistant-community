#!/usr/bin/env python3
"""Build the NCSC CAF 4.0 IGP questionnaire workbook from the official PDF.

Unlike `ncsc-caf_v4_framework_builder.py`, which keeps only the "Achieved"
statements and turns each of them into an assessable requirement, this builder
keeps the IGP tables whole: the contributing outcome is the assessable
requirement, and every indicator of good practice - Not Achieved, Partially
Achieved, Achieved - becomes a true/false question on it.

The answers decide the outcome status on their own, through the framework's
`result_aggregation: tiered_all` rule, which reproduces what the tables print:

    Not Achieved       at least one of the following statements is true
    Partially Achieved all the following statements are true
    Achieved           all the following statements are true

Each statement states its tier through the `compute_result` of its True
choice - see ANSWERS_ROWS.

PDF structure notes are the same as the sibling builder: objective / principle /
contributing-outcome headings are left-column text, and IGP tables use either
2 columns (Not Achieved / Achieved) or 3 (with Partially Achieved in between).
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

# The sibling builder owns the PDF plumbing (block iteration, noise filtering,
# statement splitting). Its filename is not a valid module name, so load it by
# path rather than duplicating ~150 lines of parsing here.
_SIBLING = Path(__file__).with_name("ncsc-caf_v4_framework_builder.py")
_spec = importlib.util.spec_from_file_location(
    "ncsc_caf_v4_framework_builder", _SIBLING
)
if _spec is None or _spec.loader is None:  # pragma: no cover - defensive
    raise SystemExit(f"Cannot load sibling builder at {_SIBLING}")
caf = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = caf  # dataclasses in the sibling need this
_spec.loader.exec_module(caf)


FRAMEWORK_SLUG = "ncsc-caf-4.0-igp"
PDF_FILENAME = "NCSC-Cyber-Assessment-Framework-4.0.pdf"
OUTPUT_FILENAME = f"{FRAMEWORK_SLUG}.xlsx"

# The source this library was generated from, downloaded 2026-09-21 from
# https://www.ncsc.gov.uk/files/NCSC-Cyber-Assessment-Framework-4.0.pdf
# (same file as the /sites/default/files/documents/ path the collection page
# links). NCSC re-issues the v4.0 PDF in place - it was last modified
# 2026-06-23 - so a hash mismatch means the source moved under us and the
# output needs re-reviewing, not that the run is wrong.
PDF_SHA256 = "1d735b2b1cae1aadd52756ede5935ccb3824417680482dbe0846b33a1e7fa3ec"

NCSC_COLLECTION_URL = "https://www.ncsc.gov.uk/collection/cyber-assessment-framework"
FRAMEWORK_NAME = "NCSC - Cyber Assessment Framework (CAF) v4.0 - IGP questionnaire"
FRAMEWORK_DESCRIPTION = (
    "National Cyber Security Centre - Cyber Assessment Framework, assessed the way "
    "the framework is designed to be used: the contributing outcome is the unit of "
    "assessment, and every indicator of good practice is answered true or false. "
    "Those answers decide the outcome status - Achieved, Partially achieved or "
    "Not achieved - by the rule the indicator tables print.\n"
    f"{NCSC_COLLECTION_URL}"
)
COPYRIGHT = f"NCSC {NCSC_COLLECTION_URL}"
PROVIDER = "NCSC"
PACKAGER = "intuitem"
PUBLICATION_DATE = "2025-08-04"

LIBRARY_META_ROWS = [
    ("type", "library"),
    ("urn", f"urn:intuitem:risk:library:{FRAMEWORK_SLUG}"),
    ("version", "1"),
    ("locale", "en"),
    ("ref_id", FRAMEWORK_SLUG),
    ("name", FRAMEWORK_NAME),
    ("description", FRAMEWORK_DESCRIPTION),
    ("copyright", COPYRIGHT),
    ("provider", PROVIDER),
    ("packager", PACKAGER),
    ("publication_date", PUBLICATION_DATE),
]

# Seeded into every new audit. The CAF has no numeric scale and no
# certification vocabulary, so the scoring and nonconformity dimensions are
# noise on the requirement page; the indicators, the outcome status and the
# evidence behind it are the whole assessment.
#
# Hiding `status` also switches the audit from status-driven to result-driven
# progress (ComplianceAssessment.progress_mode_from_visibility): an outcome
# counts as assessed once the answers produce a verdict, which is the only
# notion of "done" this framework has.
FIELD_VISIBILITY = {
    field: {"auditor": "hidden", "respondent": "hidden"}
    for field in ("score", "documentation_score", "extended_result", "status")
}

FRAMEWORK_META_ROWS = [
    ("type", "framework"),
    ("base_urn", f"urn:intuitem:risk:req_node:{FRAMEWORK_SLUG}"),
    ("urn", f"urn:intuitem:risk:framework:{FRAMEWORK_SLUG}"),
    ("ref_id", FRAMEWORK_SLUG),
    ("name", FRAMEWORK_NAME),
    ("description", FRAMEWORK_DESCRIPTION),
    ("answers_definition", "answers"),
    ("result_aggregation", "tiered_all"),
    ("field_visibility", json.dumps(FIELD_VISIBILITY)),
]

CONTENT_HEADERS = [
    "assessable",
    "depth",
    "ref_id",
    "name",
    "description",
    "annotation",
    "questions",
    "answer",
]
CONTENT_COLUMN_WIDTHS = {
    "A": 10,
    "B": 6,
    "C": 10,
    "D": 30,
    "E": 70,
    "F": 60,
    "G": 120,
    "H": 12,
}
ANSWERS_HEADERS = ["id", "question_type", "question_choices", "compute_result"]
ANSWERS_COLUMN_WIDTHS = {"A": 10, "B": 16, "C": 110, "D": 24}

# One answer definition per IGP column: True/False either way, but the result
# a True states is what tells `tiered_all` which tier the statement belongs to.
# An Achieved statement can state `compliant`, so it is a top-tier statement; a
# Partially Achieved one states the middle tier; a Not Achieved one can only
# state `non_compliant`, which makes it blocking. False on a Not Achieved
# statement states nothing - not tripping a failure condition is not an
# achievement.
ANSWER_IDS = {
    "A": "TF-A",
    "PA": "TF-PA",
    "NA": "TF-NA",
}
ANSWERS_ROWS = [
    (ANSWER_IDS["NA"], "unique_choice", "True\nFalse", "non_compliant\n/"),
    (ANSWER_IDS["PA"], "unique_choice", "True\nFalse", "partially_compliant\n/"),
    (ANSWER_IDS["A"], "unique_choice", "True\nFalse", "compliant\nnon_compliant"),
]

ANNOTATION_THREE = (
    "CAF rule for this contributing outcome: Achieved when all of the A statements "
    "are true; Partially achieved when all of the PA statements are true; Not "
    "achieved when at least one NA statement is true, or when neither set is fully met."
)
ANNOTATION_TWO = (
    "CAF rule for this contributing outcome: Achieved when all of the A statements "
    "are true, otherwise Not achieved. This outcome has no Partially achieved level."
)

# Column x0 thresholds. The published PDF lays every IGP table out on the same
# grid: statements start at x0 ~50 (NA), ~220 (PA) and ~390 (A), except in
# 2-column tables where Achieved starts at ~305.
THREE_COLUMN_BOUNDS = (("NA", 0.0, 150.0), ("PA", 150.0, 350.0), ("A", 350.0, 1e4))
TWO_COLUMN_BOUNDS = (("NA", 0.0, 200.0), ("A", 200.0, 1e4))
# Where each column's text normally starts, used to pick the closest open
# statement when a continuation is rendered outside its own column.
COLUMN_ANCHORS = {3: {"NA": 50.0, "PA": 220.0, "A": 390.0}, 2: {"NA": 50.0, "A": 305.0}}

# Every wording the v4.0 tables use for the two instruction rows, longest first.
# A handful of tables word them slightly differently (B4.d "statements below
# are true", B5.a "Any of the following statements are true").
INSTRUCTION_PHRASES = (
    "at least one of the following statements is true",
    "all the following statements below are true",
    "all of the following statements are true",
    "all the following statements are true",
    "any of the following statements are true",
)
INSTRUCTION_RE = re.compile(
    "|".join(re.escape(phrase) for phrase in INSTRUCTION_PHRASES), re.IGNORECASE
)

EXPECTED_OBJECTIVES = 4
EXPECTED_PRINCIPLES = 14
EXPECTED_OUTCOMES = 41
EXPECTED_STATEMENTS = {"NA": 184, "PA": 148, "A": 225}

# The Achieved column is also shipped, statement by statement, by the sibling
# builder. Comparing against it is the regression test for this extraction.
SHIPPED_LIBRARY = (
    Path(__file__).resolve().parents[4] / "backend/library/libraries/ncsc-caf-4.0.yaml"
)
# ref_id -> (text in the shipped library, text in the PDF). The library has a
# typo here; the PDF reads "a clear understanding".
KNOWN_LIBRARY_DIFFS = {
    "D1.a.A.1": (
        (
            "Your incident response plan is based on a clear and understanding of "
            "the security risks to network and information systems supporting your "
            "essential function(s)."
        ),
        (
            "Your incident response plan is based on a clear understanding of the "
            "security risks to network and information systems supporting your "
            "essential function(s)."
        ),
    )
}


@dataclass
class Outcome:
    ref_id: str
    name: str
    description: str = ""
    columns: int = 0
    segments: dict[str, list] = field(
        default_factory=lambda: {"NA": [], "PA": [], "A": []}
    )
    statements: dict[str, list[str]] = field(
        default_factory=lambda: {"NA": [], "PA": [], "A": []}
    )
    # The statement each column is currently building, so a continuation can be
    # matched to the column it actually belongs to rather than to where the PDF
    # happened to draw it.
    open_text: dict[str, str] = field(default_factory=dict)


@dataclass
class Node:
    depth: int
    ref_id: str
    name: str
    description: str = ""
    outcome: Outcome | None = None


def default_pdf_path() -> Path:
    return Path.cwd() / PDF_FILENAME


def default_output_path() -> Path:
    return Path.cwd() / OUTPUT_FILENAME


def column_of(x0: float, columns: int) -> str:
    bounds = THREE_COLUMN_BOUNDS if columns == 3 else TWO_COLUMN_BOUNDS
    for name, low, high in bounds:
        if low <= x0 < high:
            return name
    return bounds[-1][0]


def statement_is_open(text: str) -> bool:
    """Whether a statement is still waiting for its continuation.

    Unfinished punctuation is the usual sign, but several statements break
    across a page inside a parenthetical ("... reviewed recently (e.g." +
    "within the last 12 months)."), which ends on a full stop while clearly
    being incomplete - hence the bracket count.
    """
    return not text.rstrip().endswith((".", "!", "?")) or text.count("(") > text.count(
        ")"
    )


def owning_column(outcome: Outcome, column: str, body: str, x0: float) -> str:
    """Re-file a continuation that the PDF drew under the wrong column.

    On the A4.b page break the continuation of a Partially Achieved statement
    is rendered at the Not Achieved x-position. Geometry alone therefore
    misfiles it, corrupting two statements: the fragment is glued to whatever
    the other column said last, and its real statement is left truncated.

    A continuation whose own column has nothing open belongs to a column that
    does; where several are open, the nearest one to where the fragment was
    drawn wins.
    """
    if not body[:1].islower():
        return column

    own = outcome.open_text.get(column)
    if own and statement_is_open(own):
        return column

    candidates = [
        candidate
        for candidate, text in outcome.open_text.items()
        if candidate != column and statement_is_open(text)
    ]
    if not candidates:
        return column

    anchors = COLUMN_ANCHORS[outcome.columns]
    target = min(candidates, key=lambda c: abs(anchors[c] - x0))
    print(
        f"note: {outcome.ref_id}: continuation drawn at x0={x0:.0f} "
        f"({column} column) re-filed under {target}: {body[:60]}..."
    )
    return target


def strip_instructions(text: str) -> str:
    """Drop the "all the following statements are true" row.

    It comes through as one block per column, as a single block spanning the
    whole table (in which case it lands in the Not Achieved column), or split
    mid-phrase across two blocks - so whole phrases are removed first, and
    whatever remains is dropped when it is only a fragment of one of them.
    """
    remainder = caf.one_line(INSTRUCTION_RE.sub(" ", text))
    if not remainder:
        return ""
    lowered = remainder.lower()
    if any(lowered in phrase for phrase in INSTRUCTION_PHRASES):
        return ""
    return remainder


def parse_pdf(pdf_path: Path) -> list[Node]:
    nodes: list[Node] = []
    current: Outcome | None = None
    inside_framework = False

    for block in caf.iter_pdf_blocks(pdf_path):
        text = caf.one_line(block.text)
        if text == "The Cyber Assessment Framework":
            inside_framework = True
            continue
        if not inside_framework or not text:
            continue

        objective = caf.OBJECTIVE_RE.match(text)
        if objective:
            nodes.append(
                Node(1, objective.group(1).upper(), objective.group(2).strip())
            )
            current = None
            continue

        principle = caf.PRINCIPLE_RE.match(text)
        if principle:
            nodes.append(
                Node(2, principle.group(1).upper(), principle.group(2).strip())
            )
            current = None
            continue

        outcome_match = caf.OUTCOME_RE.match(text)
        if outcome_match and block.x0 <= caf.LEFT_TEXT_MAX_X:
            current = Outcome(outcome_match.group(1), outcome_match.group(2).strip())
            nodes.append(Node(3, current.ref_id, current.name, outcome=current))
            continue

        if current is None:
            # Objective and principle descriptions: attach to the last node.
            if nodes and block.x0 <= caf.LEFT_TEXT_MAX_X:
                nodes[-1].description = f"{nodes[-1].description} {text}".strip()
            continue

        if caf.TABLE_HEADER_RE.search(block.text):
            current.columns = 3 if "Partially Achieved" in text else 2
            continue

        if not current.columns:
            # A long outcome title wraps onto the next block (C1.f). The sibling
            # builder's rule applies: while the title has an unclosed bracket,
            # the continuation belongs to the title, not the description.
            if caf.has_unclosed_parenthesis(current.name):
                current.name, text = caf.merge_heading_continuation(current.name, text)
                if not text:
                    continue
            current.description = f"{current.description} {text}".strip()
            continue

        body = strip_instructions(text)
        if not body or caf.is_table_instruction(body):
            continue

        column = owning_column(
            current, column_of(block.x0, current.columns), body, block.x0
        )
        current.segments[column].append(
            caf.AchievedSegment(
                page=block.page, x0=block.x0, y0=block.y0, y1=block.y1, text=body
            )
        )
        # Mirror how split_statements will merge these blocks, so the next
        # continuation can be matched against the right open statement.
        previous = current.open_text.get(column, "")
        current.open_text[column] = (
            f"{previous} {body}"
            if previous and caf.is_statement_continuation(previous, body)
            else body
        )

    for node in nodes:
        if node.outcome is None:
            continue
        # Title and description are finalised on the Outcome while its table is
        # being read (bracket continuation, multi-block descriptions).
        node.name = node.outcome.name
        node.description = node.outcome.description
        for column, segments in node.outcome.segments.items():
            node.outcome.statements[column] = caf.split_statements(segments)

    return nodes


def validate(nodes: list[Node]) -> None:
    counts = {
        "objectives": sum(1 for n in nodes if n.depth == 1),
        "principles": sum(1 for n in nodes if n.depth == 2),
        "outcomes": sum(1 for n in nodes if n.depth == 3),
    }
    expected = {
        "objectives": EXPECTED_OBJECTIVES,
        "principles": EXPECTED_PRINCIPLES,
        "outcomes": EXPECTED_OUTCOMES,
    }
    statements = {
        column: sum(
            len(n.outcome.statements[column]) for n in nodes if n.outcome is not None
        )
        for column in ("NA", "PA", "A")
    }

    mismatches = [
        f"{name}: got {counts[name]}, expected {want}"
        for name, want in expected.items()
        if counts[name] != want
    ]
    mismatches += [
        f"{column} statements: got {statements[column]}, expected {want}"
        for column, want in EXPECTED_STATEMENTS.items()
        if statements[column] != want
    ]

    for node in nodes:
        outcome = node.outcome
        if outcome is None:
            continue
        if outcome.columns not in (2, 3):
            mismatches.append(f"{outcome.ref_id}: no IGP table found")
            continue
        if not outcome.statements["A"]:
            mismatches.append(f"{outcome.ref_id}: no Achieved statements")
        if outcome.columns == 3 and not outcome.statements["PA"]:
            mismatches.append(
                f"{outcome.ref_id}: Partially Achieved column with no statements"
            )
        if outcome.columns == 2 and outcome.statements["PA"]:
            mismatches.append(
                f"{outcome.ref_id}: Partially Achieved statements in a 2-column table"
            )

    if mismatches:
        raise ValueError("Unexpected parse result: " + "; ".join(mismatches))


def verify_against_library(nodes: list[Node], library_path: Path) -> int:
    """Compare the extraction with the shipped CAF 4.0 library.

    The sibling builder ships the same objective, principle and contributing
    outcome tree, plus every Achieved statement, so this is the regression test
    for this extraction. Returns the number of unexpected differences.
    """
    import yaml

    library = yaml.safe_load(library_path.read_text(encoding="utf-8"))
    shipped_nodes = library["objects"]["framework"]["requirement_nodes"]
    shipped_statements: dict[str, str] = {
        node["ref_id"]: caf.one_line(node["description"])
        for node in shipped_nodes
        if node["depth"] == 4
    }
    shipped_tree: dict[str, tuple[str, str]] = {
        node["ref_id"]: (
            caf.one_line(node.get("name") or ""),
            caf.one_line(node.get("description") or ""),
        )
        for node in shipped_nodes
        if node["depth"] < 4
    }

    unexpected = 0
    for node in nodes:
        want_tree = shipped_tree.pop(node.ref_id, None)
        got_tree = (caf.one_line(node.name), caf.one_line(node.description))
        if want_tree is not None and want_tree != got_tree:
            unexpected += 1
            print(f"  {node.ref_id}\n    library: {want_tree}\n    pdf    : {got_tree}")

        if node.outcome is None:
            continue
        for index, statement in enumerate(node.outcome.statements["A"], start=1):
            ref_id = f"{node.outcome.ref_id}.{index}"
            want = shipped_statements.pop(ref_id, None)
            if want == statement:
                continue
            known = KNOWN_LIBRARY_DIFFS.get(f"{node.outcome.ref_id}.A.{index}")
            if known and (want, statement) == known:
                continue
            unexpected += 1
            print(f"  {ref_id}\n    library: {want}\n    pdf    : {statement}")

    for ref_id, want in list(shipped_tree.items()) + list(shipped_statements.items()):
        unexpected += 1
        print(f"  {ref_id}\n    library: {want}\n    pdf    : <missing>")

    return unexpected


def build_questions(outcome: Outcome) -> tuple[str, str]:
    """Return the newline-joined question texts and their answer ids."""
    questions: list[str] = []
    answers: list[str] = []

    for column in ("NA", "PA", "A"):
        for index, statement in enumerate(outcome.statements[column], start=1):
            questions.append(f"[{outcome.ref_id}.{column}.{index}] {statement}")
            answers.append(ANSWER_IDS[column])

    return "\n".join(questions), "\n".join(answers)


def write_meta_sheet(sheet: Worksheet, rows: list[tuple[str, str]]) -> None:
    for key, value in rows:
        sheet.append([key, value])
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


def write_content_sheet(sheet: Worksheet, nodes: list[Node]) -> None:
    sheet.append(CONTENT_HEADERS)
    for node in nodes:
        if node.outcome is None:
            sheet.append(
                [
                    None,
                    node.depth,
                    node.ref_id,
                    node.name,
                    node.description,
                    None,
                    None,
                    None,
                ]
            )
            continue

        questions, answers = build_questions(node.outcome)
        annotation = ANNOTATION_THREE if node.outcome.columns == 3 else ANNOTATION_TWO
        sheet.append(
            [
                "x",
                node.depth,
                node.ref_id,
                node.name,
                node.description,
                annotation,
                questions,
                answers,
            ]
        )

    for column, width in CONTENT_COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_answers_sheet(sheet: Worksheet) -> None:
    sheet.append(ANSWERS_HEADERS)
    for row in ANSWERS_ROWS:
        sheet.append(list(row))
    for column, width in ANSWERS_COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(nodes: list[Node]) -> Workbook:
    workbook = Workbook()
    library_sheet = workbook.active
    library_sheet.title = "library_meta"
    write_meta_sheet(library_sheet, LIBRARY_META_ROWS)

    write_meta_sheet(workbook.create_sheet("framework_meta"), FRAMEWORK_META_ROWS)
    write_content_sheet(workbook.create_sheet("framework_content"), nodes)
    write_meta_sheet(
        workbook.create_sheet("answers_meta"),
        [("type", "answers"), ("name", "answers")],
    )
    write_answers_sheet(workbook.create_sheet("answers_content"))
    return workbook


def dump(nodes: list[Node]) -> None:
    for node in nodes:
        if node.outcome is None:
            print(f"{'  ' * (node.depth - 1)}{node.ref_id} {node.name}")
            continue
        outcome = node.outcome
        print(f"\n=== {outcome.ref_id} {outcome.name} ({outcome.columns} columns)")
        for column in ("NA", "PA", "A"):
            for index, statement in enumerate(outcome.statements[column], start=1):
                print(f"  {outcome.ref_id}.{column}.{index}: {statement}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", type=Path, default=default_pdf_path())
    parser.add_argument("--output", type=Path, default=default_output_path())
    parser.add_argument(
        "--dump", action="store_true", help="print the extracted statements and exit"
    )
    args = parser.parse_args()

    if not args.pdf.exists():
        raise SystemExit(f"PDF not found: {args.pdf}")

    digest = hashlib.sha256(args.pdf.read_bytes()).hexdigest()
    if digest != PDF_SHA256:
        print(f"warning: {args.pdf.name} is not the pinned source ({digest})")

    nodes = parse_pdf(args.pdf)
    validate(nodes)

    if args.dump:
        dump(nodes)
        return 0

    if SHIPPED_LIBRARY.exists():
        unexpected = verify_against_library(nodes, SHIPPED_LIBRARY)
        if unexpected:
            raise SystemExit(f"{unexpected} difference(s) from {SHIPPED_LIBRARY.name}")
        print(f"Tree and Achieved statements match {SHIPPED_LIBRARY.name}")

    build_workbook(nodes).save(args.output)
    questions = sum(
        len(n.outcome.statements["NA"])
        + len(n.outcome.statements["PA"])
        + len(n.outcome.statements["A"])
        for n in nodes
        if n.outcome is not None
    )
    print(f"Wrote {args.output} ({EXPECTED_OUTCOMES} outcomes, {questions} questions)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
