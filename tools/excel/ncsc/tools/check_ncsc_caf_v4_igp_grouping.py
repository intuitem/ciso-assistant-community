#!/usr/bin/env python3
"""Validate the grouped CAF IGP workbook and its generated YAML library."""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path

import openpyxl
import yaml


EXPECTED_OUTCOMES = 41
EXPECTED_QUESTIONS = 557
EXPECTED_GROUP_COUNTS = {2: 9, 3: 32}
EXPECTED_SLUG = "ncsc-caf-4.0-igp-grouped"


def lines(value) -> list[str]:
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def rows_by_header(sheet) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(value or "").strip() for value in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:] if any(row)]


def validate(workbook_path: Path, yaml_path: Path) -> dict:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "library_meta",
        "framework_meta",
        "framework_content",
        "answers_meta",
        "answers_content",
    ]

    library_meta = dict(
        workbook["library_meta"].iter_rows(min_col=1, max_col=2, values_only=True)
    )
    framework_meta = dict(
        workbook["framework_meta"].iter_rows(min_col=1, max_col=2, values_only=True)
    )
    assert library_meta["ref_id"] == EXPECTED_SLUG
    assert library_meta["urn"] == f"urn:intuitem:risk:library:{EXPECTED_SLUG}"
    assert framework_meta["ref_id"] == EXPECTED_SLUG
    assert framework_meta["urn"] == f"urn:intuitem:risk:framework:{EXPECTED_SLUG}"
    assert framework_meta["base_urn"] == (
        f"urn:intuitem:risk:req_node:{EXPECTED_SLUG}"
    )

    answer_rows = rows_by_header(workbook["answers_content"])
    descriptions = {
        row["id"]: row["group_description"] for row in answer_rows
    }
    assert descriptions == {
        "TF-NA": "Not Achieved",
        "TF-PA": "Partially Achieved",
        "TF-A": "Achieved",
    }

    grouped_rows = {}
    question_count = 0
    group_counts = Counter()
    for row in rows_by_header(workbook["framework_content"]):
        questions = lines(row.get("questions"))
        if not questions:
            assert not row.get("answer_group_order")
            continue
        answer_ids = lines(row.get("answer"))
        group_order = lines(row.get("answer_group_order"))
        assert len(questions) == len(answer_ids), row["ref_id"]
        expected_order = list(dict.fromkeys(answer_ids))
        assert group_order == expected_order, row["ref_id"]
        assert len(group_order) == len(set(group_order)), row["ref_id"]
        assert all(answer_id in descriptions for answer_id in group_order)
        grouped_rows[row["ref_id"]] = {
            "answer_ids": answer_ids,
            "group_order": group_order,
        }
        question_count += len(questions)
        group_counts[len(group_order)] += 1

    assert len(grouped_rows) == EXPECTED_OUTCOMES
    assert question_count == EXPECTED_QUESTIONS
    assert dict(group_counts) == EXPECTED_GROUP_COUNTS

    library = yaml.safe_load(yaml_path.read_text(encoding="utf-8"))
    assert library["urn"] == f"urn:intuitem:risk:library:{EXPECTED_SLUG}"
    framework = library["objects"]["framework"]
    assert framework["urn"] == f"urn:intuitem:risk:framework:{EXPECTED_SLUG}"

    yaml_outcomes = 0
    yaml_questions = 0
    for node in framework["requirement_nodes"]:
        questions = node.get("questions") or {}
        if not questions:
            assert not node.get("questions_properties")
            continue
        yaml_outcomes += 1
        yaml_questions += len(questions)
        source = grouped_rows[node["ref_id"]]
        groups = node["questions_properties"]["groups"]
        ordered_groups = [groups[key] for key in sorted(groups)]
        assert [group["description"] for group in ordered_groups] == [
            descriptions[answer_id] for answer_id in source["group_order"]
        ]

        question_urns = list(questions)
        expected_group_urns = {
            answer_id: [
                urn
                for urn, question_answer_id in zip(
                    question_urns, source["answer_ids"]
                )
                if question_answer_id == answer_id
            ]
            for answer_id in source["group_order"]
        }
        for answer_id, group in zip(source["group_order"], ordered_groups):
            assert group["order"] == expected_group_urns[answer_id], node["ref_id"]
        flattened = [urn for group in ordered_groups for urn in group["order"]]
        assert flattened == question_urns, node["ref_id"]
        assert len(flattened) == len(set(flattened)), node["ref_id"]

    assert yaml_outcomes == EXPECTED_OUTCOMES
    assert yaml_questions == EXPECTED_QUESTIONS
    return {
        "outcomes": yaml_outcomes,
        "questions": yaml_questions,
        "group_counts": dict(sorted(group_counts.items())),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    script_dir = Path(__file__).resolve().parent
    ncsc_dir = script_dir.parent
    repository = script_dir.parents[3]
    parser.add_argument(
        "--workbook",
        type=Path,
        default=ncsc_dir / f"{EXPECTED_SLUG}.xlsx",
    )
    parser.add_argument(
        "--yaml",
        type=Path,
        default=repository
        / "backend"
        / "library"
        / "libraries"
        / f"{EXPECTED_SLUG}.yaml",
    )
    args = parser.parse_args()
    result = validate(args.workbook, args.yaml)
    print(
        "Validated grouped CAF IGP: "
        f"{result['outcomes']} outcomes, {result['questions']} questions, "
        f"groups {result['group_counts']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
