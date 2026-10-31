#!/usr/bin/env python3
"""Build a CISO Assistant OWASP Top 10 workbook from normalized JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


DEFAULT_LOCALE = "en"
CONTENT_HEADERS = ["ref_id", "name", "description", "annotation"]


def as_text(value: Any) -> str:
    return "" if value is None else str(value)


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as source:
        payload = json.load(source)
    if not isinstance(payload, dict):
        raise ValueError("The normalized JSON root must be an object")
    return payload


def write_key_value_sheet(
    workbook: Workbook, title: str, rows: list[tuple[str, str]]
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = True
    for key, value in rows:
        worksheet.append([as_text(key), as_text(value)])

    for row in worksheet.iter_rows():
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 120


def write_content_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    rows: list[dict[str, str]],
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = "A2"
    worksheet.append(headers)

    for row in rows:
        worksheet.append([as_text(row.get(header, "")) for header in headers])

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, header in enumerate(headers, start=1):
        if header == "ref_id":
            width = 14
        elif header.startswith("name"):
            width = 42
        elif header.startswith("description"):
            width = 100
        else:
            width = 120
        worksheet.column_dimensions[get_column_letter(index)].width = width


def build_excel_from_json(input_path: Path, output_path: Path) -> dict[str, Any]:
    payload = load_json(input_path)
    year = int(payload["year"])
    default_locale = str(payload.get("default_locale", DEFAULT_LOCALE))
    locales = [str(locale) for locale in payload["exported_locales"]]
    if not locales or locales[0] != default_locale:
        raise ValueError("The default locale must be the first exported locale")

    library_translations = payload["library"]["translations"]
    default_library = library_translations[default_locale]
    library_suffix = f"owasp-top-10-web_{year}"
    library_urn = f"urn:intuitem:risk:library:{library_suffix}"
    library_rows: list[tuple[str, str]] = [
        ("type", "library"),
        ("urn", library_urn),
        ("version", "1"),
        ("locale", default_locale),
        ("ref_id", f"OWASP-top-10-Web_{year}"),
        ("name", as_text(default_library["name"])),
        ("description", as_text(default_library["description"])),
        ("copyright", as_text(payload["library"]["copyright"])),
        ("provider", "OWASP"),
        ("packager", "intuitem"),
    ]

    for locale in locales[1:]:
        translation = library_translations[locale]
        library_rows.extend(
            [
                (f"name[{locale}]", as_text(translation["name"])),
                (f"description[{locale}]", as_text(translation["description"])),
            ]
        )

    base_urn = (
        "urn:intuitem:risk:threat"
        if year == 2021
        else f"urn:intuitem:risk:{library_suffix}"
    )
    threat_meta_rows = [("type", "threats"), ("base_urn", base_urn)]

    headers = CONTENT_HEADERS.copy()
    for locale in locales[1:]:
        headers.extend(
            [
                f"name[{locale}]",
                f"description[{locale}]",
                f"annotation[{locale}]",
            ]
        )

    content_rows: list[dict[str, str]] = []
    for threat in payload["threats"]:
        translations = threat["translations"]
        default_translation = translations[default_locale]
        row = {
            "ref_id": as_text(threat["ref_id"]),
            "name": as_text(default_translation["name"]),
            "description": as_text(default_translation["description"]),
            "annotation": as_text(default_translation["annotation"]),
        }
        for locale in locales[1:]:
            translation = translations[locale]
            row[f"name[{locale}]"] = as_text(translation["name"])
            row[f"description[{locale}]"] = as_text(translation["description"])
            row[f"annotation[{locale}]"] = as_text(translation["annotation"])
        content_rows.append(row)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = f"OWASP Top 10 Web - {year}"
    workbook.properties.subject = "CISO Assistant threat library"
    workbook.properties.creator = "intuitem"

    write_key_value_sheet(workbook, "library_meta", library_rows)
    write_key_value_sheet(workbook, "threats_meta", threat_meta_rows)
    write_content_sheet(workbook, "threats_content", headers, content_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {
        "year": year,
        "languages": locales,
        "threats": len(content_rows),
        "base_urn": base_urn,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a CISO Assistant OWASP Top 10 XLSX file from normalized JSON"
    )
    parser.add_argument("input", help="Normalized JSON input path")
    parser.add_argument("--output", required=True, help="Output XLSX path")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    stats = build_excel_from_json(Path(args.input), Path(args.output))
    print(f"Excel generated: {args.output}")
    print(f"Threats exported: {stats['threats']}")
    print(f"Languages exported: {len(stats['languages'])}")
    print(f"Languages: {', '.join(stats['languages'])}")


if __name__ == "__main__":
    main()
