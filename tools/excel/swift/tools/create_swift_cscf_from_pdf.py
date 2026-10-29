"""Build a CISO Assistant Swift CSCF workbook from a Swift CSCF PDF.

The PDF is the only framework input.  The generated workbook intentionally does
not read, copy, or otherwise depend on a prior Excel workbook.

Requirements:
    pip install PyMuPDF openpyxl

Usage:
    python create_swift_cscf_from_pdf.py CSCF_v2026_202507015.pdf 2026
    python create_swift_cscf_from_pdf.py CSCF_v2027_20260710.pdf 2027
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DESCRIPTION = (
    "This document establishes a set of mandatory and advisory security controls "
    "for the operating environment of Swift users. Mandatory security controls "
    "build on existing guidance and establish a security baseline for the entire "
    "user community. Advisory controls are optional best practices that Swift "
    "recommends each user to implement in the operating environment. This document "
    "must be read in conjunction with the [CSP FAQ](https://www2.swift.com/kb/#/tip/5021823) "
    "Knowledge Base article 5021823 which provides additional valuable information."
)

ARCHITECTURES = ("A1", "A2", "A3", "A4", "B")
ARCHITECTURE_CENTERS = {
    "A1": 372.0,
    "A2": 407.0,
    "A3": 442.0,
    "A4": 477.0,
    "B": 512.0,
}
PRIMARY_CENTER = 220.0
CONTROL_PATTERN = re.compile(r"^\d+(?:\.\d+)?A?$")
DATE_PATTERN = re.compile(
    r"\b\d{1,2} (?:January|February|March|April|May|June|July|August|"
    r"September|October|November|December) 20\d{2}\b"
)


@dataclass(frozen=True)
class SummaryRow:
    """One principle or security control from the PDF summary table."""

    ref_id: str
    label: str
    depth: int
    architectures: tuple[str, ...]
    grey_columns: frozenset[str]


@dataclass(frozen=True)
class Subsection:
    """An assessable lettered implementation-guideline subsection."""

    ref_id: str
    label: str


def is_grey(fill: tuple[float, float, float] | None) -> bool:
    """Return whether a PDF drawing fill is the neutral advisory grey."""
    if fill is None:
        return False
    red, green, blue = fill
    return 0.75 <= red <= 0.95 and max(fill) - min(fill) <= 0.02


def find_summary_page(document: pymupdf.Document) -> pymupdf.Page:
    """Find the page containing the mandatory/advisory controls summary table."""
    for page in document:
        text = page.get_text("text")
        if (
            "Security Controls Summary Table" in text
            and "Mandatory and Advisory Security Controls" in text
            and "Architecture Type" in text
        ):
            return page
    raise ValueError("The PDF does not contain a Security Controls Summary Table.")


def extract_issue_date(document: pymupdf.Document) -> str:
    """Extract the date displayed on the last line of the PDF's first page."""
    matches = DATE_PATTERN.findall(document[0].get_text("text"))
    if not matches:
        raise ValueError("No issue date was found on the first PDF page.")
    return matches[-1]


def table_reference_words(page: pymupdf.Page) -> list[tuple[str, float, float, float]]:
    """Return numbered table rows, ordered from top to bottom."""
    rows: list[tuple[str, float, float, float]] = []
    for x0, y0, x1, _y1, text, *_rest in page.get_text("words"):
        # The table is above the PDF footer.  The date in that footer can otherwise
        # be mistaken for a top-level principle number.
        if y0 < 100 or y0 > 740 or x0 > 145 or not CONTROL_PATTERN.fullmatch(text):
            continue
        if "." not in text and x0 > 103:
            continue
        if "." in text and x0 < 103:
            continue
        rows.append((text, x0, x1, y0))

    if not rows:
        raise ValueError("No framework rows were found in the summary table.")
    return sorted(rows, key=lambda row: row[3])


def row_label(
    words: Iterable[tuple[float, float, float, float, str, int, int, int]],
    row_y: float,
    reference_end: float,
) -> str:
    """Read the text in the left-most cell of a summary-table row."""
    label_words = [
        word
        for word in words
        if row_y - 1 <= word[1] <= row_y + 2
        and reference_end < word[0] < 354
    ]
    label = " ".join(word[4] for word in sorted(label_words, key=lambda word: word[0]))
    if not label:
        raise ValueError(f"Could not read the label for the row at y={row_y:.1f}.")
    return label


def grey_columns_for_row(
    grey_rectangles: Iterable[pymupdf.Rect], row_y: float
) -> frozenset[str]:
    """Find shaded cells that overlap a table row."""
    columns = {"primary": PRIMARY_CENTER, **ARCHITECTURE_CENTERS}
    result = set()
    for column, center_x in columns.items():
        if any(
            rectangle.x0 <= center_x <= rectangle.x1
            and rectangle.y0 - 1 <= row_y <= rectangle.y1 + 1
            for rectangle in grey_rectangles
        ):
            result.add(column)
    return frozenset(result)


def architectures_for_row(
    words: Iterable[tuple[float, float, float, float, str, int, int, int]],
    row_y: float,
) -> tuple[str, ...]:
    """Map the table's bullet points to architecture columns."""
    selected = []
    for x0, y0, x1, _y1, text, *_rest in words:
        if text != "●" or abs(y0 - row_y) > 5:
            continue
        center_x = (x0 + x1) / 2
        architecture = min(
            ARCHITECTURE_CENTERS,
            key=lambda candidate: abs(ARCHITECTURE_CENTERS[candidate] - center_x),
        )
        if abs(ARCHITECTURE_CENTERS[architecture] - center_x) < 12:
            selected.append(architecture)
    return tuple(architecture for architecture in ARCHITECTURES if architecture in selected)


def extract_summary_rows(page: pymupdf.Page) -> list[SummaryRow]:
    """Extract framework nodes and architecture applicability from the summary."""
    words = page.get_text("words")
    references = table_reference_words(page)
    grey_rectangles = [
        drawing["rect"]
        for drawing in page.get_drawings()
        if is_grey(drawing.get("fill"))
    ]

    rows = []
    for ref_id, _reference_start, reference_end, row_y in references:
        depth = 2 if "." in ref_id else 1
        rows.append(
            SummaryRow(
                ref_id=ref_id,
                label=row_label(words, row_y, reference_end),
                depth=depth,
                architectures=architectures_for_row(words, row_y),
                grey_columns=grey_columns_for_row(grey_rectangles, row_y),
            )
        )

    controls = [row for row in rows if row.depth == 2]
    if not controls:
        raise ValueError("No security controls were found in the summary table.")
    return rows


def page_lines(page: pymupdf.Page) -> list[tuple[float, float, str]]:
    """Return PDF text lines with their left and top coordinates."""
    grouped_words: dict[tuple[int, int], list[tuple]] = {}
    for word in page.get_text("words"):
        grouped_words.setdefault((word[5], word[6]), []).append(word)

    lines = []
    for words in grouped_words.values():
        ordered = sorted(words, key=lambda word: word[0])
        lines.append(
            (
                min(word[0] for word in ordered),
                min(word[1] for word in ordered),
                " ".join(word[4] for word in ordered),
            )
        )
    return sorted(lines, key=lambda line: (line[1], line[0]))


def extract_subsections(
    document: pymupdf.Document, summary_rows: list[SummaryRow]
) -> dict[str, list[Subsection]]:
    """Extract lettered implementation-guideline headings from detailed controls.

    Swift lays out these headings with ``a)``, ``b)``, and so on in the left
    margin.  The x-coordinate distinguishes them from explanatory bullet lists
    and nested examples such as ``d.1`` and ``o``.
    """
    known_controls = {row.ref_id for row in summary_rows if row.depth == 2}
    result: dict[str, list[Subsection]] = {ref_id: [] for ref_id in known_controls}
    current_control: str | None = None
    subsection_pattern = re.compile(r"^([a-z])\)(?:\s+(.+))?$")

    for page in document:
        lines = page_lines(page)
        for line_index, (x0, y0, text) in enumerate(lines):
            control_match = re.match(r"^(\d+\.\d+A?)\s+", text)
            if (
                control_match
                and x0 <= 80
                and control_match.group(1) in known_controls
            ):
                current_control = control_match.group(1)
                continue

            subsection_match = subsection_pattern.match(text)
            if not subsection_match or current_control is None or not 65 <= x0 <= 78:
                continue
            letter, label = subsection_match.groups()
            if label is None:
                # A long heading can be emitted as a separate PDF text block on
                # the same baseline, immediately to the right of ``c)``.
                right_hand_lines = [
                    candidate_text
                    for candidate_x, candidate_y, candidate_text in lines[line_index + 1 :]
                    if abs(candidate_y - y0) <= 1 and candidate_x > x0
                ]
                if right_hand_lines:
                    label = right_hand_lines[0]
            if not label:
                raise ValueError(
                    f"Could not read the subsection heading for {current_control}.{letter}."
                )
            result[current_control].append(
                Subsection(ref_id=f"{current_control}.{letter}", label=label.strip())
            )

    for control_ref, subsections in result.items():
        if not subsections:
            continue
        letters = [subsection.ref_id.rsplit(".", 1)[1] for subsection in subsections]
        expected = [chr(ord("a") + index) for index in range(len(letters))]
        if letters != expected:
            raise ValueError(
                f"Unexpected subsection sequence for {control_ref}: {', '.join(letters)}"
            )
    return result


def implementation_groups(row: SummaryRow) -> tuple[str, ...]:
    """Translate summary-table shading and architecture columns into CISO groups."""
    if row.depth != 2:
        return ()

    if row.ref_id.endswith("A"):
        if "primary" not in row.grey_columns:
            raise ValueError(
                f"{row.ref_id} is advisory but its complete row is not shaded in the PDF."
            )
        return ("A",)

    groups = []
    for architecture in row.architectures:
        if architecture in row.grey_columns:
            if architecture == "A4":
                groups.append("AA4")
            elif architecture == "B":
                groups.append("AB")
            else:
                raise ValueError(
                    f"{row.ref_id} has a shaded {architecture} cell, for which no IG exists."
                )
        else:
            groups.append(architecture)

    if not groups:
        raise ValueError(f"{row.ref_id} has no applicable architecture in the summary table.")
    return tuple(groups)


def apply_sheet_style(sheet, widths: list[float]) -> None:
    """Apply a compact, readable style without relying on a template workbook."""
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = True


def append_metadata(sheet, values: list[tuple[str, str]]) -> None:
    for key, value in values:
        sheet.append((key, value))
    sheet.column_dimensions["A"].width = 35
    sheet.column_dimensions["B"].width = 115
    sheet.sheet_view.showGridLines = True


def build_workbook(
    rows: list[SummaryRow],
    subsections_by_control: dict[str, list[Subsection]],
    year: str,
    issue_date: str,
) -> Workbook:
    """Create the complete CISO Assistant workbook without loading a template."""
    framework_id = f"swift-cscf-v{year}"
    title = (
        f"Swift Customer Security Controls Framework (CSCF) v{year} "
        f"[{issue_date}] - Outline"
    )

    workbook = Workbook()
    library_meta = workbook.active
    library_meta.title = "library_meta"
    append_metadata(
        library_meta,
        [
            ("type", "library"),
            ("urn", f"urn:intuitem:risk:library:{framework_id}"),
            ("version", "1"),
            ("locale", "en"),
            ("ref_id", framework_id),
            ("name", title),
            ("description", DESCRIPTION),
            ("copyright", "SWIFT"),
            ("provider", "SWIFT"),
            ("packager", "intuitem"),
        ],
    )

    measures_meta = workbook.create_sheet("mesures_meta")
    append_metadata(
        measures_meta,
        [
            ("type", "framework"),
            ("base_urn", f"urn:intuitem:risk:req_node:{framework_id}"),
            ("urn", f"urn:intuitem:risk:framework:{framework_id}"),
            ("ref_id", framework_id),
            ("name", title),
            ("description", DESCRIPTION),
            ("implementation_groups_definition", "imp_grp"),
        ],
    )

    measures_content = workbook.create_sheet("mesures_content")
    measures_content.append(
        ("assessable", "depth", "ref_id", "name", "description", "implementation_groups")
    )
    for row in rows:
        if row.depth == 1:
            measures_content.append((None, 1, row.ref_id, row.label, None, None))
        elif subsections := subsections_by_control[row.ref_id]:
            measures_content.append((None, 2, row.ref_id, None, row.label, None))
            groups = ",".join(implementation_groups(row))
            for subsection in subsections:
                measures_content.append(
                    ("x", 3, subsection.ref_id, None, subsection.label, groups)
                )
        else:
            measures_content.append(
                ("x", 2, row.ref_id, None, row.label, ",".join(implementation_groups(row)))
            )
    apply_sheet_style(measures_content, [12, 9, 12, 58, 72, 30])

    implementation_meta = workbook.create_sheet("imp_grp_meta")
    append_metadata(
        implementation_meta,
        [("type", "implementation_groups"), ("name", "imp_grp")],
    )

    implementation_content = workbook.create_sheet("imp_grp_content")
    implementation_content.append(("ref_id", "name", "description"))
    for ref_id, name in (
        ("A", "Advisory Controls"),
        ("A1", "Architecture A1"),
        ("A2", "Architecture A2"),
        ("A3", "Architecture A3"),
        ("A4", "Architecture A4"),
        ("B", "Architecture B"),
        ("AB", "Architecture B (Advisory)"),
        ("AA4", "Architecture A4 (Advisory)"),
    ):
        implementation_content.append((ref_id, name, None))
    apply_sheet_style(implementation_content, [12, 38, 55])

    return workbook


def convert(pdf_path: Path, year: str, output_path: Path) -> None:
    """Read one PDF and save its standalone workbook."""
    if not pdf_path.is_file():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")
    if not re.fullmatch(r"20\d{2}", year):
        raise ValueError("The year must have four digits, for example 2026.")

    with pymupdf.open(pdf_path) as document:
        issue_date = extract_issue_date(document)
        rows = extract_summary_rows(find_summary_page(document))
        subsections_by_control = extract_subsections(document, rows)

    workbook = build_workbook(rows, subsections_by_control, year, issue_date)
    workbook.save(output_path)

    controls = sum(row.depth == 2 for row in rows)
    subsections = sum(len(items) for items in subsections_by_control.values())
    print(f"Created {output_path}")
    print(
        f"Extracted {controls} controls, {subsections} subsections, "
        f"and {len(rows) - controls} principles."
    )


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf", type=Path, help="Source Swift CSCF PDF")
    parser.add_argument("year", help="CSCF year, for example 2026")
    parser.add_argument(
        "--output",
        type=Path,
        help="Optional output XLSX path (default: next to the PDF)",
    )
    return parser.parse_args()


def main() -> int:
    arguments = parse_arguments()
    output = arguments.output or arguments.pdf.with_name(f"swift-cscf-v{arguments.year}.xlsx")
    try:
        convert(arguments.pdf, arguments.year, output)
    except (FileNotFoundError, ValueError, pymupdf.FileDataError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
