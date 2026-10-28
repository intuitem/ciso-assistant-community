import io

import pytest
from django.urls import reverse
from knox.models import AuthToken
from openpyxl import load_workbook
from rest_framework.test import APIClient

from core.apps import startup
from iam.models import Folder, User, UserGroup
from privacy.models import Processing
from tprm.models import Entity


@pytest.fixture
def admin_client():
    startup(sender=None)
    admin = User.objects.create_superuser("admin@processor-register-tests.com")
    admin_group = UserGroup.objects.get(name="BI-UG-ADM")
    admin.folder = admin_group.folder
    admin.save()
    admin_group.user_set.add(admin)
    client = APIClient()
    token = AuthToken.objects.create(user=admin)
    client.credentials(HTTP_AUTHORIZATION=f"Token {token[1]}")
    return client


@pytest.mark.django_db
def test_processor_register_filter_and_export(admin_client):
    folder = Folder.objects.create(
        parent_folder=Folder.get_root_folder(), name="register folder"
    )
    client_a = Entity.objects.create(name="Client A", folder=folder)
    client_b = Entity.objects.create(name="Client B", folder=folder)
    hosting = Processing.objects.create(
        name="Hosting", folder=folder, role="privacy_data_processor"
    )
    hosting.controllers.set([client_a, client_b])
    payroll = Processing.objects.create(name="Payroll", folder=folder)
    assert payroll.role == "privacy_controller"

    url = reverse("processings-list")
    res = admin_client.get(url, {"role": "privacy_data_processor"})
    assert res.status_code == 200
    assert [p["name"] for p in res.json()["results"]] == ["Hosting"]
    res = admin_client.get(url, {"controllers": str(client_a.id)})
    assert [p["name"] for p in res.json()["results"]] == ["Hosting"]

    res = admin_client.get(
        reverse("processings-export-record-xlsx", kwargs={"pk": str(hosting.id)})
    )
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert "Security measures" in wb.sheetnames
    controllers = [
        row[0] for row in wb["Controllers"].iter_rows(min_row=2, values_only=True)
    ]
    assert sorted(controllers) == ["Client A", "Client B"]
    processing_sheet = list(wb["Processing"].iter_rows(values_only=True))
    role_col = processing_sheet[0].index("role")
    assert processing_sheet[1][role_col] == "Data Processor"


@pytest.mark.django_db
def test_controllers_only_kept_for_non_controller_roles(admin_client):
    folder = Folder.objects.create(
        parent_folder=Folder.get_root_folder(), name="role rule folder"
    )
    partner = Entity.objects.create(name="Partner", folder=folder)
    processing = Processing.objects.create(
        name="Shared CRM", folder=folder, role="privacy_data_processor"
    )
    processing.controllers.set([partner])
    url = reverse("processings-detail", kwargs={"pk": str(processing.id)})

    res = admin_client.patch(url, {"role": "privacy_joint_controller"}, format="json")
    assert res.status_code == 200
    assert list(processing.controllers.all()) == [partner]

    res = admin_client.patch(url, {"role": "privacy_controller"}, format="json")
    assert res.status_code == 200
    assert not processing.controllers.exists()

    res = admin_client.patch(url, {"controllers": [str(partner.id)]}, format="json")
    assert res.status_code == 200
    assert not processing.controllers.exists()
