import io
from datetime import date

import openpyxl
import pytest

from core.models import (
    AppliedControl,
    Asset,
    ComplianceAssessment,
    Framework,
    Perimeter,
    RequirementAssessment,
    RequirementNode,
    RiskAssessment,
    RiskScenario,
    Terminology,
    Threat,
)
from ebios_rm.models import (
    EbiosRMStudy,
    ElementaryAction,
    FearedEvent,
    KillChain,
    Stakeholder,
)
from ebios_rm.tests.test_kill_chain_steps import _operating_mode
from rest_framework.test import APIClient

from core.utils import UserGroupCodename
from iam.models import Folder, User, UserGroup
from tprm.models import Entity


@pytest.mark.django_db
def test_export_xlsx_sanitizes_audit_sheet_name(
    admin_client, basic_ebios_rm_study_fixture
):
    study = basic_ebios_rm_study_fixture
    folder = Folder.get_root_folder()
    fw = Framework.objects.create(
        name="FW export", urn="urn:test:fw-export", folder=folder
    )
    rn = RequirementNode.objects.create(
        urn="urn:test:rn-export", framework=fw, assessable=True, folder=folder
    )
    ca = ComplianceAssessment.objects.create(
        name="Socle ISO 27001:2022 [v1]/*?\\ – Nu",
        framework=fw,
        folder=folder,
        perimeter=Perimeter.objects.create(name="p-export", folder=folder),
    )
    RequirementAssessment.objects.create(
        compliance_assessment=ca, requirement=rn, folder=folder
    )
    study.compliance_assessments.add(ca)

    response = admin_client.get(f"/api/ebios-rm/studies/{study.id}/export-xlsx/")

    assert response.status_code == 200
    sheets = openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames
    assert "1.4.1 Socle ISO 27001_2022 _v1_" in sheets


@pytest.mark.django_db
def test_export_import_round_trips_workshop_5(
    admin_client, basic_ebios_rm_study_fixture
):
    study = basic_ebios_rm_study_fixture
    root = Folder.get_root_folder()
    ra = RiskAssessment.objects.create(
        name="WS5 risk assessment",
        ref_id="RA.1",
        version="2",
        folder=root,
        risk_matrix=study.risk_matrix,
        ebios_rm_study=study,
    )
    measure = AppliedControl.objects.create(
        name="MFA rollout",
        ref_id="M.1",
        folder=root,
        status="in_progress",
        eta=date(2027, 3, 31),
    )
    existing = AppliedControl.objects.create(name="Firewall", folder=root)
    scenario = RiskScenario.objects.create(
        risk_assessment=ra,
        name="Ransomware via phishing",
        ref_id="R.1",
        current_proba=2,
        current_impact=3,
        residual_proba=0,
        residual_impact=1,
        treatment="mitigate",
        justification="MFA cuts the entry vector",
    )
    scenario.applied_controls.set([measure])
    scenario.existing_applied_controls.set([existing])

    export = admin_client.get(f"/api/ebios-rm/studies/{study.id}/export-xlsx/")
    assert export.status_code == 200

    target = Folder.objects.create(name="Imported", parent_folder=root)
    response = admin_client.post(
        "/api/data-wizard/load-file/",
        data=export.content,
        content_type="application/octet-stream",
        HTTP_X_MODEL_TYPE="EbiosRMStudyExcel",
        HTTP_X_FOLDER_ID=str(target.id),
        HTTP_X_MATRIX_ID=str(study.risk_matrix.id),
        HTTP_CONTENT_DISPOSITION="attachment; filename=study.xlsx",
    )
    assert response.status_code == 200, response.content
    results = response.json()["results"]
    assert results["successful"] == 1, results["errors"]

    imported = EbiosRMStudy.objects.get(folder=target)
    imported_ra = imported.last_risk_assessment
    assert (imported_ra.name, imported_ra.ref_id, imported_ra.version) == (
        "WS5 risk assessment",
        "RA.1",
        "2",
    )
    assert imported_ra.folder == target
    rs = imported_ra.risk_scenarios.get()
    assert (rs.name, rs.ref_id, rs.treatment, rs.justification) == (
        "Ransomware via phishing",
        "R.1",
        "mitigate",
        "MFA cuts the entry vector",
    )
    assert (rs.current_proba, rs.current_impact) == (2, 3)
    assert (rs.residual_proba, rs.residual_impact) == (0, 1)
    assert [c.name for c in rs.existing_applied_controls.all()] == ["Firewall"]
    imported_measure = rs.applied_controls.get()
    assert imported_measure.folder == target
    assert (imported_measure.ref_id, imported_measure.status) == ("M.1", "in_progress")
    assert imported_measure.eta == date(2027, 3, 31)


@pytest.mark.django_db
def test_export_xlsx_hides_linked_objects_the_user_cannot_view(
    admin_client, ebios_rm_matrix_fixture
):
    domain = Folder.objects.create(
        name="Study domain",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
        create_iam_groups=True,
    )
    Folder.create_default_ug_and_ra(domain)
    # Global-folder objects are visible from every domain, a sibling's are not
    other = Folder.objects.create(
        name="Other domain",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )
    study = EbiosRMStudy.objects.create(
        name="Scoped study", risk_matrix=ebios_rm_matrix_fixture, folder=domain
    )
    visible_asset = Asset.objects.create(name="Visible asset", folder=domain)
    hidden_asset = Asset.objects.create(name="Hidden asset", folder=other)
    hidden_asset.parent_assets.add(visible_asset)
    visible_asset.parent_assets.add(hidden_asset)
    study.assets.set([visible_asset, hidden_asset])
    feared_event = FearedEvent.objects.create(name="Feared event", ebios_rm_study=study)
    feared_event.assets.set([visible_asset, hidden_asset])

    framework = Framework.objects.create(
        name="FW scoped", urn="urn:test:fw-scoped", folder=other
    )
    rn = RequirementNode.objects.create(
        urn="urn:test:rn-scoped", framework=framework, assessable=True, folder=other
    )
    hidden_audit = ComplianceAssessment.objects.create(
        name="Hidden audit", framework=framework, folder=other
    )
    RequirementAssessment.objects.create(
        compliance_assessment=hidden_audit, requirement=rn, folder=other
    )
    study.compliance_assessments.add(hidden_audit)

    category, _ = Terminology.objects.get_or_create(
        name="partner",
        field_path=Terminology.FieldPath.ENTITY_RELATIONSHIP,
        defaults={"is_visible": True},
    )
    stakeholder = Stakeholder.objects.create(
        ebios_rm_study=study,
        entity=Entity.objects.create(name="Hidden entity", folder=other),
        category=category,
    )
    visible = AppliedControl.objects.create(name="Visible control", folder=domain)
    hidden = AppliedControl.objects.create(name="Hidden control", folder=other)
    stakeholder.applied_controls.set([visible, hidden])

    operating_mode = _operating_mode(study)
    operating_mode.operational_scenario.attack_path.stakeholders.add(stakeholder)
    step = KillChain.objects.create(
        operating_mode=operating_mode,
        elementary_action=ElementaryAction.objects.create(
            name="Hidden elementary action", folder=other
        ),
    )
    step.assets.set([hidden_asset])
    KillChain.objects.create(
        operating_mode=operating_mode,
        elementary_action=ElementaryAction.objects.create(
            name="Visible elementary action", folder=domain
        ),
    ).assets.set([hidden_asset])

    ra = RiskAssessment.objects.create(
        name="Scoped RA",
        folder=domain,
        risk_matrix=ebios_rm_matrix_fixture,
        ebios_rm_study=study,
    )
    scenario = RiskScenario.objects.create(risk_assessment=ra, name="Scenario")
    scenario.applied_controls.set([visible, hidden])
    scenario.existing_applied_controls.set([hidden])
    scenario.assets.set([visible_asset, hidden_asset])
    scenario.threats.set([Threat.objects.create(name="Hidden threat", folder=other)])

    reader = User.objects.create_user("reader@ebios-export-tests.com")
    UserGroup.objects.get(
        folder=domain, name=str(UserGroupCodename.READER)
    ).user_set.add(reader)
    client = APIClient()
    client.force_authenticate(reader)

    response = client.get(f"/api/ebios-rm/studies/{study.id}/export-xlsx/")

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    cells = [
        str(cell)
        for ws in workbook.worksheets
        for row in ws.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    ]
    for name in (
        "Visible asset",
        "Visible control",
        "Visible elementary action",
        "Scenario",
    ):
        assert any(name in c for c in cells), name
    assert [c for c in cells if "Hidden" in c] == []
    assert not any(sheet.startswith("1.4") for sheet in workbook.sheetnames)
