"""CCB CyberFundamentals (CyFun®): framework URNs, assurance levels and the
official self-assessment tools audits are exported to (the "cyfun-xlsx"
framework export)."""

import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from functools import partial
from pathlib import Path

from openpyxl import load_workbook

from .framework_exports import ExportFile, FrameworkExport, register

CYFUN_2023_URN = "urn:intuitem:risk:framework:ccb-cff-2023-03-01"
CYFUN_2025_URN = "urn:intuitem:risk:framework:ccb-cyfun2025"

LEVELS = ("basic", "important", "essential")

# Implementation groups of each level, with their key measure (xK) and, in
# CyFun 2025, management aspect (xG) subsets. An audit without groups covers
# ESSENTIAL.
GROUP_LEVELS = {
    group: level
    for level, prefix in zip(LEVELS, "BIE")
    for group in (prefix, f"{prefix}K", f"{prefix}G")
}

# The official CyFun 2025 tool of each level (core/templates/core) and the
# column of its requirement text; the documentation score, implementation
# score and comment columns follow at +1, +2 and +7. BASIC has no "Assurance
# level" column.
TEMPLATES = {
    "basic": ("CyFun2025_Self-Assessment_tool_BASIC_v2026_02_20.xlsx", 5),
    "important": ("CyFun2025_Self-Assessment_tool_IMPORTANT_v2026_02_20.xlsx", 6),
    "essential": ("CyFun2025_Self-Assessment_tool_ESSENTIAL_v3.1.xlsx", 6),
}
FUNCTION_SHEETS_2025 = ("GOVERN", "IDENTIFY", "PROTECT", "DETECT", "RESPOND", "RECOVER")

# The CyFun 2023 tool covers the three levels, one "<LEVEL> Details" sheet each,
# with the requirement in E, the scores in G and H and comments in M.
TEMPLATE_2023 = "CyFun2023_Self_Assessment_tool_V20251021.xlsx"

_REF_ID = re.compile(r"^([A-Z]{2}\.[A-Z]{2})-(\d+)(?:[.-](\d+))?$")
# CyFun 2023 ids keep the CCB's own forms, including "R.AC-3.4". The tool
# doesn't always put a colon after them ("PR.AC-3.1 The organisation's ...").
_REF_ID_2023 = re.compile(r"^((?:BASIC_|IMPORTANT_)?[A-Z]{1,2}\.[A-Z]{2}-\d+\.\d+)\b")


def level_for_groups(selected_implementation_groups) -> str:
    """Assurance level of a CyFun audit: the highest of its groups."""
    levels = {GROUP_LEVELS.get(group) for group in selected_implementation_groups or []}
    levels.discard(None)
    return max(levels, key=LEVELS.index) if levels else "essential"


def normalize_ref_id(text) -> str | None:
    """CyFun 2025 ref_id of a tool's requirement cell ("GV.OC-01.1: ..."), or None.

    The tools don't always zero-pad the subcategory (ID.AM-5.1) and sometimes
    put a hyphen before the requirement number (DE.CM-03-1).
    """
    match = _REF_ID.match(str(text or "").split(":")[0].strip().rstrip("."))
    if not match:
        return None
    prefix, subcategory, requirement = match.groups()
    ref_id = f"{prefix}-{int(subcategory):02d}"
    return f"{ref_id}.{requirement}" if requirement else ref_id


def ref_id_2023(text, level) -> str | None:
    """CyFun 2023 ref_id of a requirement cell of the tool's sheet for *level*.

    A sheet prefixes the requirements carried over from lower levels (BASIC_,
    IMPORTANT_); its own are bare, which the library prefixes too, except for
    ESSENTIAL.
    """
    match = _REF_ID_2023.match(str(text or "").strip().upper())
    if not match:
        return None
    ref_id = match.group(1)
    if level != "essential" and not ref_id.startswith(("BASIC_", "IMPORTANT_")):
        ref_id = f"{level.upper()}_{ref_id}"
    return ref_id


@dataclass(frozen=True)
class ExportTemplate:
    """Where an official self-assessment tool takes each requirement."""

    file: str  # in core/templates/core
    sheets: tuple[str, ...]
    requirement_column: int
    doc_column: int
    impl_column: int
    comment_column: int
    ref_id: Callable[[object], str | None]  # library ref_id of a requirement cell


def export_template(
    framework_urn, selected_implementation_groups
) -> ExportTemplate | None:
    """The tool an audit exports to: its framework version's, at its level."""
    level = level_for_groups(selected_implementation_groups)
    if framework_urn == CYFUN_2025_URN:
        file, column = TEMPLATES[level]
        return ExportTemplate(
            file,
            FUNCTION_SHEETS_2025,
            column,
            column + 1,
            column + 2,
            column + 7,
            normalize_ref_id,
        )
    if framework_urn == CYFUN_2023_URN:
        return ExportTemplate(
            TEMPLATE_2023,
            (f"{level.upper()} Details",),
            5,
            7,
            8,
            13,
            partial(ref_id_2023, level=level),
        )
    return None


def _supports(audit) -> bool:
    return (
        export_template(audit.framework.urn, audit.selected_implementation_groups)
        is not None
    )


def build_self_assessment(audit) -> ExportFile:
    """The official tool of the audit's CyFun version and assurance level,
    filled with its scores and observations. The tool lists only that level's
    requirements and applies its own N/A score and thresholds."""
    from .models import RequirementAssessment
    from .utils import escape_excel_formula, sanitize_xlsx_value

    template = export_template(
        audit.framework.urn, audit.selected_implementation_groups
    )
    wb = load_workbook(
        Path(__file__).resolve().parent / "templates" / "core" / template.file
    )

    # ref_id -> (sheet, row) of every requirement the tool lists
    rows = {}
    for sheet_name in template.sheets:
        ws = wb[sheet_name]
        for row in range(1, ws.max_row + 1):
            cell_value = ws.cell(row=row, column=template.requirement_column).value
            if cell_value and isinstance(cell_value, str):
                ref_id = template.ref_id(cell_value)
                if ref_id:
                    rows[ref_id] = (ws, row)

    requirement_assessments = (
        RequirementAssessment.objects.filter(compliance_assessment=audit)
        .select_related("requirement")
        .filter(requirement__assessable=True)
    )
    for ra in requirement_assessments:
        target = rows.get((ra.requirement.ref_id or "").upper())
        if target is None:
            continue
        ws, row = target
        if ra.result == RequirementAssessment.Result.NOT_APPLICABLE:
            ws.cell(row=row, column=template.doc_column, value="N/A")
            ws.cell(row=row, column=template.impl_column, value="N/A")
        elif ra.is_scored:
            # Only the scores the audit counts.
            if audit.show_documentation_score and ra.documentation_score is not None:
                ws.cell(
                    row=row,
                    column=template.doc_column,
                    value=ra.documentation_score,
                )
            if ra.score is not None:
                ws.cell(row=row, column=template.impl_column, value=ra.score)
        if ra.observation:
            ws.cell(
                row=row,
                column=template.comment_column,
                value=sanitize_xlsx_value(escape_excel_formula(ra.observation)),
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    return ExportFile(
        buffer.getvalue(),
        f"{audit.name}_CyFun_Self-Assessment.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


register(
    FrameworkExport(
        ref_id="cyfun-xlsx",
        title="exportCyFunAssessment",
        description="exportCyFunAssessmentDesc",
        format="XLSX",
        supports=_supports,
        build=build_self_assessment,
    )
)
