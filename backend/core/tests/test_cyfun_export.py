import io

import pytest
from django.urls import reverse
from knox.models import AuthToken
from openpyxl import load_workbook
from rest_framework.test import APIClient

from core import cyfun
from core.apps import startup
from core.models import (
    ComplianceAssessment,
    Framework,
    Perimeter,
    RequirementAssessment,
    RequirementNode,
)
from core.utils import EVERYONE_EDIT
from iam.models import Folder, User, UserGroup

SCORES_VISIBLE = {
    field: EVERYONE_EDIT for field in ("score", "is_scored", "documentation_score")
}


@pytest.fixture
def admin_client():
    startup(sender=None)
    admin = User.objects.create_superuser("admin@cyfun-export-tests.com")
    admin_group = UserGroup.objects.get(name="BI-UG-ADM")
    admin.folder = admin_group.folder
    admin.save()
    admin_group.user_set.add(admin)
    client = APIClient()
    token = AuthToken.objects.create(user=admin)
    client.credentials(HTTP_AUTHORIZATION=f"Token {token[1]}")
    return client


@pytest.fixture
def cyfun_audit():
    """A CyFun 2025 audit on three BASIC requirements, two of which the BASIC
    tool lists with irregular ids (ID.AM-5.1, DE.CM-03-1)."""
    root = Folder.get_root_folder()
    folder = Folder.objects.create(parent_folder=root, name="cyfun export folder")
    perimeter = Perimeter.objects.create(name="cyfun export perimeter", folder=folder)
    framework = Framework.objects.create(
        name="CyFun 2025",
        urn=cyfun.CYFUN_2025_URN,
        min_score=1,
        max_score=5,
        folder=root,
        implementation_groups_definition=[
            {"ref_id": ref_id, "name": ref_id} for ref_id in ("B", "I", "E")
        ],
    )
    ca = ComplianceAssessment.objects.create(
        field_visibility=SCORES_VISIBLE,
        name="CyFun export",
        framework=framework,
        folder=folder,
        perimeter=perimeter,
    )
    for ref_id, doc, impl, result in (
        ("GV.OC-03.1", 2, 3, RequirementAssessment.Result.PARTIALLY_COMPLIANT),
        ("ID.AM-05.1", 4, 5, RequirementAssessment.Result.COMPLIANT),
        ("DE.CM-03.1", None, None, RequirementAssessment.Result.NOT_APPLICABLE),
    ):
        node = RequirementNode.objects.create(
            urn=f"urn:test:cyfun-export:{ref_id.lower()}",
            ref_id=ref_id,
            framework=framework,
            assessable=True,
            folder=root,
        )
        RequirementAssessment.objects.create(
            compliance_assessment=ca,
            requirement=node,
            folder=folder,
            is_scored=impl is not None,
            score=impl,
            documentation_score=doc,
            result=result,
        )
    return ca


def _export_url(ca, export_id="cyfun-xlsx"):
    return reverse(
        "compliance-assessments-framework-export",
        kwargs={"pk": str(ca.pk), "export_id": export_id},
    )


def _export(client, ca):
    response = client.get(_export_url(ca))
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.content))


def _scores(workbook, sheet, ref_id, requirement_column):
    """(documentation, implementation) written on the row of ref_id."""
    for row in workbook[sheet].iter_rows(min_row=3):
        if cyfun.normalize_ref_id(row[requirement_column - 1].value) == ref_id:
            return row[requirement_column].value, row[requirement_column + 1].value
    raise AssertionError(f"{ref_id} not found in {sheet}")


@pytest.mark.django_db
class TestCyfunExportPerLevel:
    def test_writes_only_the_scores_the_audit_counts(self, admin_client, cyfun_audit):
        """Scores kept while unscored, and hidden documentation scores, are not
        exported: the tool keeps its own minimum."""
        cyfun_audit.selected_implementation_groups = ["B"]
        cyfun_audit.save()
        RequirementAssessment.objects.filter(
            compliance_assessment=cyfun_audit, requirement__ref_id="ID.AM-05.1"
        ).update(is_scored=False)
        workbook = _export(admin_client, cyfun_audit)
        assert _scores(workbook, "IDENTIFY", "ID.AM-05.1", 5) == (1, 1)

        cyfun_audit.show_documentation_score = False
        cyfun_audit.save()
        workbook = _export(admin_client, cyfun_audit)
        assert _scores(workbook, "GOVERN", "GV.OC-03.1", 5) == (1, 3)

    def test_observations_with_control_characters_are_exported(
        self, admin_client, cyfun_audit
    ):
        RequirementAssessment.objects.filter(
            compliance_assessment=cyfun_audit, requirement__ref_id="GV.OC-03.1"
        ).update(observation="pasted\x0bfrom Word")
        _export(admin_client, cyfun_audit)

    @pytest.mark.parametrize(
        "groups, level, summary",
        [
            (["B"], "basic", "BASIC Summary"),
            (["I"], "important", "IMPORTANT Summary"),
            ([], "essential", "ESSENTIAL Summary"),
        ],
    )
    def test_uses_the_tool_of_the_audit_level(
        self, admin_client, cyfun_audit, groups, level, summary
    ):
        cyfun_audit.selected_implementation_groups = groups
        cyfun_audit.save()

        workbook = _export(admin_client, cyfun_audit)

        assert summary in workbook.sheetnames
        _, requirement_column = cyfun.TEMPLATES[level]
        assert _scores(workbook, "GOVERN", "GV.OC-03.1", requirement_column) == (2, 3)
        assert _scores(workbook, "IDENTIFY", "ID.AM-05.1", requirement_column) == (
            4,
            5,
        )
        assert _scores(workbook, "DETECT", "DE.CM-03.1", requirement_column) == (
            "N/A",
            "N/A",
        )


@pytest.fixture
def cyfun2023_audit():
    """A CyFun 2023 audit with a requirement of each level, keeping the CCB's
    ids: BASIC_ and IMPORTANT_ prefixes, and the tool's own "R.AC-3.4"."""
    root = Folder.get_root_folder()
    folder = Folder.objects.create(parent_folder=root, name="cyfun 2023 export folder")
    perimeter = Perimeter.objects.create(name="cyfun 2023 perimeter", folder=folder)
    framework = Framework.objects.create(
        name="CyFun 2023",
        urn=cyfun.CYFUN_2023_URN,
        min_score=1,
        max_score=5,
        folder=root,
        implementation_groups_definition=[
            {"ref_id": ref_id, "name": ref_id} for ref_id in ("B", "I", "E")
        ],
    )
    ca = ComplianceAssessment.objects.create(
        field_visibility=SCORES_VISIBLE,
        name="CyFun 2023 export",
        framework=framework,
        folder=folder,
        perimeter=perimeter,
    )
    for ref_id, doc, impl, result in (
        ("BASIC_ID.AM-1.1", 2, 3, RequirementAssessment.Result.PARTIALLY_COMPLIANT),
        ("IMPORTANT_ID.AM-1.2", 4, 4, RequirementAssessment.Result.COMPLIANT),
        ("ID.AM-1.4", 3, 2, RequirementAssessment.Result.PARTIALLY_COMPLIANT),
        ("R.AC-3.4", None, None, RequirementAssessment.Result.NOT_APPLICABLE),
    ):
        node = RequirementNode.objects.create(
            urn=f"urn:test:cyfun2023-export:{ref_id.lower()}",
            ref_id=ref_id,
            framework=framework,
            assessable=True,
            folder=root,
        )
        RequirementAssessment.objects.create(
            compliance_assessment=ca,
            requirement=node,
            folder=folder,
            is_scored=impl is not None,
            score=impl,
            documentation_score=doc,
            result=result,
        )
    return ca


def _scores_2023(workbook, level, ref_id):
    """(documentation, implementation) on the row of ref_id in the level's sheet."""
    for row in workbook[f"{level.upper()} Details"].iter_rows(min_row=3):
        if cyfun.ref_id_2023(row[4].value, level) == ref_id:
            return row[6].value, row[7].value
    raise AssertionError(f"{ref_id} not found in {level} sheet")


@pytest.mark.django_db
class TestCyfun2023Export:
    def test_basic_audit_fills_the_basic_sheet(self, admin_client, cyfun2023_audit):
        cyfun2023_audit.selected_implementation_groups = ["B"]
        cyfun2023_audit.save()
        workbook = _export(admin_client, cyfun2023_audit)
        assert _scores_2023(workbook, "basic", "BASIC_ID.AM-1.1") == (2, 3)

    def test_important_audit_fills_the_important_sheet(
        self, admin_client, cyfun2023_audit
    ):
        cyfun2023_audit.selected_implementation_groups = ["I"]
        cyfun2023_audit.save()
        workbook = _export(admin_client, cyfun2023_audit)
        # BASIC requirements appear prefixed in the IMPORTANT sheet.
        assert _scores_2023(workbook, "important", "BASIC_ID.AM-1.1") == (2, 3)
        assert _scores_2023(workbook, "important", "IMPORTANT_ID.AM-1.2") == (4, 4)

    def test_essential_audit_fills_the_essential_sheet(
        self, admin_client, cyfun2023_audit
    ):
        workbook = _export(admin_client, cyfun2023_audit)
        assert _scores_2023(workbook, "essential", "BASIC_ID.AM-1.1") == (2, 3)
        assert _scores_2023(workbook, "essential", "IMPORTANT_ID.AM-1.2") == (4, 4)
        assert _scores_2023(workbook, "essential", "ID.AM-1.4") == (3, 2)
        assert _scores_2023(workbook, "essential", "R.AC-3.4") == ("N/A", "N/A")


CYFUN_EXPORT = {
    "ref_id": "cyfun-xlsx",
    "title": "exportCyFunAssessment",
    "description": "exportCyFunAssessmentDesc",
    "format": "XLSX",
}


@pytest.mark.django_db
class TestFrameworkExports:
    def _detail(self, client, ca):
        url = reverse("compliance-assessments-detail", kwargs={"pk": str(ca.pk)})
        return client.get(url).json()

    def test_cyfun_audits_offer_the_self_assessment_export(
        self, admin_client, cyfun_audit, cyfun2023_audit
    ):
        assert self._detail(admin_client, cyfun_audit)["framework_exports"] == [
            CYFUN_EXPORT
        ]
        assert self._detail(admin_client, cyfun2023_audit)["framework_exports"] == [
            CYFUN_EXPORT
        ]

    def test_other_frameworks_offer_none(self, admin_client, cyfun_audit):
        cyfun_audit.framework.urn = "urn:test:framework:other"
        cyfun_audit.framework.save()
        assert self._detail(admin_client, cyfun_audit)["framework_exports"] == []
        assert admin_client.get(_export_url(cyfun_audit)).status_code == 400

    def test_unknown_export_is_not_found(self, admin_client, cyfun_audit):
        response = admin_client.get(_export_url(cyfun_audit, "no-such-export"))
        assert response.status_code == 404
