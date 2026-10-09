"""fill_template writes cells into an Excel template and leaves every other part
of the file as it is."""

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core.xlsx_template import fill_template

MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@pytest.fixture
def template(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scores"
    ws["A1"] = 1
    ws["A1"].font = Font(bold=True)
    ws["B1"] = "=A1*2"
    ws["A3"] = 3
    path = tmp_path / "template.xlsx"
    wb.save(path)
    return path


def _sheet(content):
    return load_workbook(io.BytesIO(content))["Scores"]


def test_writes_numbers_and_texts_keeping_the_style(template):
    content = fill_template(template, {"Scores": {(1, 1): 4, (1, 3): "N/A"}})
    ws = _sheet(content)
    assert ws["A1"].value == 4
    assert ws["A1"].font.bold
    assert ws["C1"].value == "N/A"
    # The formula is kept.
    assert ws["B1"].value == "=A1*2"


def test_creates_missing_rows_in_order(template):
    content = fill_template(template, {"Scores": {(2, 1): "  spaced  "}})
    ws = _sheet(content)
    assert ws["A2"].value == "  spaced  "
    assert ws["A3"].value == 3


def test_untouched_parts_are_copied_as_is(template):
    content = fill_template(template, {"Scores": {(1, 1): 4}})
    with (
        zipfile.ZipFile(template) as before,
        zipfile.ZipFile(io.BytesIO(content)) as after,
    ):
        assert before.namelist() == after.namelist()
        assert before.read("xl/styles.xml") == after.read("xl/styles.xml")
        assert b'fullCalcOnLoad="1"' in after.read("xl/workbook.xml")


def test_formula_cells_are_refused(template):
    with pytest.raises(ValueError, match="holds a formula"):
        fill_template(template, {"Scores": {(1, 2): 5}})


def test_unknown_sheets_are_refused(template):
    with pytest.raises(ValueError, match="Unknown template sheets"):
        fill_template(template, {"Missing": {(1, 1): 1}})
