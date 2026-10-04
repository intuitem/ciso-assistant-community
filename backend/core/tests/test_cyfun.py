from pathlib import Path
from types import SimpleNamespace

import pytest
import yaml
from openpyxl import load_workbook

from core import cyfun
from core.models import Framework
from core.utils import EVERYONE_EDIT, HIDDEN, build_initial_field_visibility

BACKEND = Path(__file__).resolve().parents[2]
FUNCTION_SHEETS = ("GOVERN", "IDENTIFY", "PROTECT", "DETECT", "RESPOND", "RECOVER")


def _framework():
    library = (BACKEND / "library/libraries/cyfun2025.yaml").read_text()
    return yaml.safe_load(library)["objects"]["framework"]


def test_cyfun_frameworks_declare_their_scoring():
    """The CCB scores CyFun on its own 1-5 scale, as a mean of category means."""
    cyfun2025 = _framework()
    assert cyfun2025["score_scale_locked"] is True
    assert cyfun2025["score_calculation_method"] == "average_of_averages"
    library = (BACKEND / "library/libraries/ccb-cff-2023-03-01.yaml").read_text()
    cyfun2023 = yaml.safe_load(library)["objects"]["framework"]
    assert cyfun2023["score_scale_locked"] is True
    assert cyfun2023["score_calculation_method"] == "average_of_averages"
    # Every control starts at the minimum in the CCB tools.
    assert cyfun2025["score_defaults_to_minimum"] is True
    assert cyfun2023["score_defaults_to_minimum"] is True


@pytest.mark.parametrize(
    "groups, target",
    [
        ([], 3),
        (["E"], 3),
        (["I", "IK"], 3),
        (["B"], 2.5),
        (["BK", "BG"], 2.5),
        # BASIC and IMPORTANT groups together make an IMPORTANT audit.
        (["B", "I"], 3),
    ],
)
@pytest.mark.django_db
def test_new_cyfun_2025_audits_count_na_as_their_level_target(groups, target):
    """The CCB tools count an N/A requirement as their level's key measure
    threshold (2.5 for BASIC, 3 above)."""
    framework = _framework()
    scoring = Framework(
        score_calculation_method=framework["score_calculation_method"],
        anchor_na_to_target=framework["anchor_na_to_target"],
        target_score=framework["target_score"],
        implementation_groups_definition=framework["implementation_groups_definition"],
    ).default_scoring_for(groups)
    assert scoring["anchor_na_to_target"] is True
    assert scoring["target_score"] == target


@pytest.mark.parametrize("library", ["cyfun2025", "ccb-cff-2023-03-01"])
def test_new_audits_show_scores_to_auditors_and_respondents(library):
    text = (BACKEND / f"library/libraries/{library}.yaml").read_text()
    framework = yaml.safe_load(text)["objects"]["framework"]
    visibility = build_initial_field_visibility(
        SimpleNamespace(field_visibility=framework["field_visibility"])
    )
    for field in ("score", "is_scored", "documentation_score"):
        assert visibility[field] == EVERYONE_EDIT


def test_new_cyfun_2025_audits_hide_extended_result_and_progress_status():
    """The CCB tools have neither; progress then follows compliance."""
    visibility = build_initial_field_visibility(
        SimpleNamespace(field_visibility=_framework()["field_visibility"])
    )
    assert visibility["extended_result"] == HIDDEN
    assert visibility["status"] == HIDDEN


@pytest.mark.parametrize(
    "groups, level",
    [
        (None, "essential"),
        ([], "essential"),
        (["B"], "basic"),
        (["BK", "BG"], "basic"),
        (["IK"], "important"),
        # BASIC and IMPORTANT groups together make an IMPORTANT audit.
        (["B", "I"], "important"),
        (["E", "B"], "essential"),
    ],
)
def test_level_for_groups(groups, level):
    assert cyfun.level_for_groups(groups) == level


@pytest.mark.parametrize(
    "text, ref_id",
    [
        ("GV.OC-01.1: The organisation's mission shall be established", "GV.OC-01.1"),
        ("PR.AA-05.4.", "PR.AA-05.4"),
        # Irregular ids found in the BASIC and IMPORTANT tools.
        ("ID.AM-5.1: Information is protected", "ID.AM-05.1"),
        ("DE.CM-03-1: Personnel activity", "DE.CM-03.1"),
        ("Requirement", None),
        (None, None),
    ],
)
def test_normalize_ref_id(text, ref_id):
    assert cyfun.normalize_ref_id(text) == ref_id


@pytest.mark.parametrize("level", cyfun.LEVELS)
def test_each_tool_lists_exactly_its_level_requirements(level):
    """Every requirement row of a tool maps onto a framework requirement of its
    level, and none of them is missing, so the export fills every row."""
    group = {"basic": "B", "important": "I", "essential": "E"}[level]
    expected = {
        node["ref_id"]
        for node in _framework()["requirement_nodes"]
        if node.get("assessable") and group in node.get("implementation_groups", [])
    }

    template, requirement_column = cyfun.TEMPLATES[level]
    workbook = load_workbook(BACKEND / "core/templates/core" / template)
    listed = [
        cyfun.normalize_ref_id(row[0])
        for sheet in FUNCTION_SHEETS
        for row in workbook[sheet].iter_rows(
            min_row=3,
            min_col=requirement_column,
            max_col=requirement_column,
            values_only=True,
        )
        if row[0]
    ]
    assert None not in listed
    assert sorted(listed) == sorted(expected)


def _marked_management_aspects(level):
    """Requirements a tool marks as linked to the management aspects (column B,
    whose marks may span merged cells)."""
    template, requirement_column = cyfun.TEMPLATES[level]
    workbook = load_workbook(BACKEND / "core/templates/core" / template)
    marked = set()
    for sheet in FUNCTION_SHEETS:
        ws = workbook[sheet]
        merged = {
            row: ws.cell(cells.min_row, 2).value
            for cells in ws.merged_cells.ranges
            if cells.min_col <= 2 <= cells.max_col
            for row in range(cells.min_row, cells.max_row + 1)
        }
        for row in range(3, ws.max_row + 1):
            mark = ws.cell(row, 2).value or merged.get(row)
            ref_id = cyfun.normalize_ref_id(ws.cell(row, requirement_column).value)
            if ref_id and mark and "management" in str(mark).lower():
                marked.add(ref_id)
    return marked


@pytest.mark.parametrize(
    "level, group, count",
    [("basic", "BG", 0), ("important", "IG", 9), ("essential", "EG", 15)],
)
def test_management_aspect_groups_match_the_tools(level, group, count):
    """The xG groups carry the CCB's current "controls linked to the management
    aspects" (called "governance measures" before the 2026-02-20 tools)."""
    tagged = {
        node["ref_id"]
        for node in _framework()["requirement_nodes"]
        if group in node.get("implementation_groups", [])
    }
    assert tagged == _marked_management_aspects(level)
    assert len(tagged) == count


def test_implementation_group_ids_are_stable():
    """Audits and campaigns store these ids: dropping one on a library update
    would silently widen their scope."""
    groups = [g["ref_id"] for g in _framework()["implementation_groups_definition"]]
    assert groups == ["B", "I", "E", "BK", "IK", "EK", "BG", "IG", "EG"]


@pytest.mark.parametrize(
    "text, level, ref_id",
    [
        ("ID.AM-1.1: Physical devices", "essential", "ID.AM-1.1"),
        ("ID.AM-1.1: Physical devices", "basic", "BASIC_ID.AM-1.1"),
        ("BASIC_ID.AM-1.1: Physical devices", "important", "BASIC_ID.AM-1.1"),
        # No colon after the id in the BASIC tool.
        ("PR.AC-3.1 The organisation's wireless", "basic", "BASIC_PR.AC-3.1"),
        ("R.AC-3.4: Remote access", "essential", "R.AC-3.4"),
        ("NO REQUIREMENT  / Guidance to be considered", "basic", None),
        (None, "basic", None),
    ],
)
def test_ref_id_2023(text, level, ref_id):
    assert cyfun.ref_id_2023(text, level) == ref_id


@pytest.mark.parametrize("level", cyfun.LEVELS)
def test_cyfun_2023_tool_lists_exactly_its_level_requirements(level):
    """Each "<LEVEL> Details" sheet of the CyFun 2023 tool maps onto the library
    requirements of that level, so the export fills every row."""
    library = (BACKEND / "library/libraries/ccb-cff-2023-03-01.yaml").read_text()
    framework = yaml.safe_load(library)["objects"]["framework"]
    group = {"basic": "B", "important": "I", "essential": "E"}[level]
    expected = {
        node["ref_id"].upper()
        for node in framework["requirement_nodes"]
        if node.get("assessable") and group in node.get("implementation_groups", [])
    }
    workbook = load_workbook(BACKEND / "core/templates/core" / cyfun.TEMPLATE_2023)
    listed = [
        cyfun.ref_id_2023(row[0], level)
        for row in workbook[f"{level.upper()} Details"].iter_rows(
            min_row=3, min_col=5, max_col=5, values_only=True
        )
        if row[0] and cyfun.ref_id_2023(row[0], level)
    ]
    assert sorted(listed) == sorted(expected)
