"""Saved import mappings: stored through the API, applied by analyze to files
with the same layout, counted when a real import uses them."""

import io
import json
from pathlib import Path

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from data_wizard.models import ImportMapping

SAMPLES = Path(__file__).parent / "samples"
TRACKER = "02_isms_controls_tracker_EN.xlsx"
TRACKER_HEADERS = [
    "Control ID",
    "Control",
    "Details",
    "Category",
    "CSF Function",
    "Status",
    "Priority",
    "Effort",
    "Owner",
    "Due date",
    "Evidence link",
]
MAPPINGS_URL = "/api/data-wizard/mappings/"
ANALYZE_URL = "/api/data-wizard/mapping/analyze/"
IMPORT_URL = "/api/data-wizard/mapping/import/"


def _tracker_sheet(**changes):
    entry = {
        "source": "Controls",
        "headers": TRACKER_HEADERS,
        "header_row": 2,
        "model_type": "AppliedControl",
        "target": "Sheet1",
        "columns": {
            "Control ID": "ref_id",
            "Control": "name",
            "Details": "description",
            "Status": "status",
            "Due date": "eta",
        },
        "values": {"Status": {"Planned": "to_do", "Not started": "to_do"}},
    }
    entry.update(changes)
    return entry


def _upload(name, data=None):
    return SimpleUploadedFile(name, data or (SAMPLES / name).read_bytes())


def _analyze(client, name, data=None, **extra):
    resp = client.post(
        ANALYZE_URL, {"file": _upload(name, data), **extra}, format="multipart"
    )
    assert resp.status_code == 200, resp.json()
    return {s["name"]: s for s in resp.json()["sheets"]}


@pytest.fixture
def saved(root_folder, admin_user):
    return ImportMapping.objects.create(
        name="ISMS tracker", folder=root_folder, sheets=[_tracker_sheet()]
    )


@pytest.mark.django_db
class TestApi:
    def test_create_records_author(self, api_client, root_folder, admin_user):
        resp = api_client.post(
            MAPPINGS_URL,
            {
                "name": "ISMS tracker",
                "folder": str(root_folder.id),
                "sheets": [_tracker_sheet()],
            },
            format="json",
        )
        assert resp.status_code == 201, resp.json()
        assert ImportMapping.objects.get(name="ISMS tracker").created_by == admin_user

    @pytest.mark.parametrize(
        "change",
        [
            {"columns": {"Not a header": "name"}},
            {"target": "Nope"},
            {"model_type": "TaskTemplate"},
            {"headers": "Control ID"},
            {"defaults": {"status": "active"}},  # status is mapped
            {"defaults": {"nope": "x"}},
        ],
    )
    def test_rejects_what_the_importer_would_reject(
        self, api_client, root_folder, change
    ):
        resp = api_client.post(
            MAPPINGS_URL,
            {
                "name": "Bad",
                "folder": str(root_folder.id),
                "sheets": [_tracker_sheet(**change)],
            },
            format="json",
        )
        assert resp.status_code == 400
        assert "sheets" in resp.json()

    def test_list(self, api_client, saved):
        names = [m["name"] for m in api_client.get(MAPPINGS_URL).json()["results"]]
        assert names == ["ISMS tracker"]


@pytest.mark.django_db
class TestAnalyzeAppliesSavedMapping:
    def test_same_layout_uses_saved_columns_and_values(self, api_client, saved):
        sheet = _analyze(api_client, TRACKER)["Controls"]
        assert sheet["saved_mapping"]["name"] == "ISMS tracker"
        columns = sheet["mapping"]["columns"]
        # Detection alone sends "Due date" to expiry_date: the saved choice wins.
        assert columns["Due date"] == {"key": "eta", "reason": "saved"}
        assert sheet["mapping"]["values"]["Status"]["Planned"] == {
            "value": "to_do",
            "reason": "saved",
        }
        # A known column left unmapped was a reviewed choice: it stays ignored.
        assert "Evidence link" not in columns

    def test_defaults_come_back(self, api_client, saved):
        saved.sheets = [_tracker_sheet(defaults={"priority": 2})]
        saved.save()
        sheet = _analyze(api_client, TRACKER)["Controls"]
        assert sheet["mapping"]["defaults"] == {"priority": 2}

    def test_extra_column_still_matches_and_is_reported(self, api_client, saved):
        wb = openpyxl.load_workbook(SAMPLES / TRACKER)
        ws = wb["Controls"]
        ws.cell(row=2, column=ws.max_column + 1, value="Budget")
        buffer = io.BytesIO()
        wb.save(buffer)
        sheet = _analyze(api_client, TRACKER, buffer.getvalue())["Controls"]
        assert sheet["saved_mapping"]["new_columns"] == ["Budget"]
        assert sheet["mapping"]["columns"]["Due date"]["key"] == "eta"

    def test_other_layout_is_not_matched(self, api_client, saved):
        sheets = _analyze(api_client, "03_vendor_register.xlsx")
        assert all(s.get("saved_mapping") is None for s in sheets.values())

    def test_can_be_turned_off(self, api_client, saved):
        sheet = _analyze(api_client, TRACKER, use_saved="false")["Controls"]
        assert sheet["saved_mapping"] is None
        assert sheet["mapping"]["columns"]["Due date"]["key"] != "eta"

    def test_user_override_wins(self, api_client, saved):
        overrides = {"Controls": {"model_type": "Policy", "sheet": "Sheet1"}}
        sheet = _analyze(api_client, TRACKER, overrides=json.dumps(overrides))[
            "Controls"
        ]
        assert sheet["saved_mapping"] is None
        assert sheet["target"]["model_type"] == "Policy"


@pytest.mark.django_db
class TestUsage:
    def _import(self, client, folder, saved, dry_run):
        entry = saved.sheets[0]
        step = {
            "model_type": entry["model_type"],
            "sheets": [
                {
                    k: entry[k]
                    for k in ("source", "target", "header_row", "columns", "values")
                }
            ],
        }
        return client.post(
            IMPORT_URL,
            {
                "file": _upload(TRACKER),
                "mapping": json.dumps({"steps": [step]}),
                "options": json.dumps(
                    {
                        "folder_id": str(folder.id),
                        "on_conflict": "skip",
                        "saved_mapping_ids": [str(saved.id)],
                    }
                ),
                "dry_run": str(dry_run).lower(),
            },
            format="multipart",
        )

    def test_dry_run_does_not_count(
        self, api_client, domain_folder, saved, all_accessible
    ):
        assert self._import(api_client, domain_folder, saved, True).status_code == 200
        saved.refresh_from_db()
        assert saved.use_count == 0

    def test_import_counts(self, api_client, domain_folder, saved, all_accessible):
        resp = self._import(api_client, domain_folder, saved, False)
        assert resp.json()["steps"][0]["results"]["created"] >= 1
        saved.refresh_from_db()
        assert saved.use_count == 1
        assert saved.last_used_at is not None
