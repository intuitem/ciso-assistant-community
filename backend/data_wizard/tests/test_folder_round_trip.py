"""A domains export re-imports as is: same columns, matched back by internal_id."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from core.models import FilteringLabel
from iam.models import Folder, UserGroup

URL = "/api/data-wizard/load-file/"


def _export(client, fmt: str = "xlsx") -> bytes:
    resp = client.get(f"/api/folders/export_{fmt}/")
    assert resp.status_code == 200, resp.content
    return resp.content


def _import(client, content: bytes, on_conflict: str, fmt: str = "xlsx"):
    return client.post(
        URL,
        data=content,
        content_type="application/octet-stream",
        HTTP_X_MODEL_TYPE="Folder",
        HTTP_X_ON_CONFLICT=on_conflict,
        HTTP_CONTENT_DISPOSITION=f"attachment; filename=domains_export.{fmt}",
    )


@pytest.fixture
def domains(root_folder):
    alpha = Folder.objects.create(name="Alpha", parent_folder=root_folder)
    alpha.filtering_labels.set([FilteringLabel.objects.create(label="Corporate")])
    beta = Folder.objects.create(name="Beta", parent_folder=root_folder)
    beta.filtering_labels.set(
        [
            FilteringLabel.objects.create(label="Finance"),
            FilteringLabel.objects.create(label="IT"),
        ]
    )
    gamma = Folder.objects.create(name="Gamma", parent_folder=root_folder)
    return alpha, beta, gamma


def _labels(folder) -> set[str]:
    return set(folder.filtering_labels.values_list("label", flat=True))


@pytest.mark.django_db
class TestFolderRoundTrip:
    def test_export_columns_match_import(self, api_client, domains, all_accessible):
        ws = load_workbook(io.BytesIO(_export(api_client))).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == (
            "internal_id",
            "name",
            "description",
            "domain",
            "labels",
            "create_iam_groups",
        )
        exported = {row[1]: row for row in rows[1:]}
        # The root folder is implicit in the import, so it is not exported.
        assert set(exported) == {"Alpha", "Beta", "Gamma"}
        assert exported["Alpha"][3] is None  # blank parent: placed at the root
        assert set(exported["Beta"][4].split(",")) == {"Finance", "IT"}

    def test_unchanged_export_reimports_without_changes(
        self, api_client, domains, all_accessible
    ):
        resp = _import(api_client, _export(api_client), "update")
        assert resp.status_code == 200, resp.json()
        results = resp.json()["results"]
        assert results["failed"] == 0, results["errors"]
        assert results["created"] == 0
        assert results["updated"] == 3

        alpha, beta, gamma = (Folder.objects.get(id=f.id) for f in domains)
        assert _labels(alpha) == {"Corporate"}
        assert _labels(beta) == {"Finance", "IT"}
        assert _labels(gamma) == set()
        assert (
            Folder.objects.filter(content_type=Folder.ContentType.DOMAIN).count() == 3
        )

    def test_edited_export_updates_and_creates(
        self, api_client, domains, all_accessible
    ):
        alpha, beta, _ = domains
        wb = load_workbook(io.BytesIO(_export(api_client)))
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=2):
            if row[1].value == "Alpha":
                row[2].value = "edited"
            if row[1].value == "Beta":
                row[4].value = "Audit"
        ws.append([None, "Delta", None, None, "New,Corporate"])
        edited = io.BytesIO()
        wb.save(edited)

        resp = _import(api_client, edited.getvalue(), "update")
        assert resp.status_code == 200, resp.json()
        results = resp.json()["results"]
        assert results["failed"] == 0, results["errors"]
        assert results["created"] == 1
        assert results["updated"] == 3

        alpha.refresh_from_db()
        assert alpha.description == "edited"
        assert _labels(beta) == {"Audit"}
        assert _labels(Folder.objects.get(name="Delta")) == {"New", "Corporate"}
        # Existing labels are reused, not duplicated.
        assert FilteringLabel.objects.filter(label="Corporate").count() == 1

    @pytest.mark.parametrize("fmt", ["xlsx", "csv"])
    def test_formula_like_values_survive_the_round_trip(
        self, api_client, root_folder, all_accessible, fmt
    ):
        """The export escapes cells starting with = + - @; the import undoes it."""
        folder = Folder.objects.create(
            name="-Ops", description="- first item", parent_folder=root_folder
        )
        folder.filtering_labels.set(
            [
                FilteringLabel.objects.create(label="-dash"),
                FilteringLabel.objects.create(label="plain"),
            ]
        )

        resp = _import(api_client, _export(api_client, fmt), "update", fmt)
        assert resp.status_code == 200, resp.json()
        results = resp.json()["results"]
        assert results["failed"] == 0, results["errors"]
        assert results["updated"] == 1

        folder.refresh_from_db()
        assert folder.name == "-Ops"
        assert folder.description == "- first item"
        assert _labels(folder) == {"-dash", "plain"}
        assert FilteringLabel.objects.count() == 2

    def test_csv_export_reimports(self, api_client, domains, all_accessible):
        resp = _import(api_client, _export(api_client, "csv"), "update", "csv")
        assert resp.status_code == 200, resp.json()
        results = resp.json()["results"]
        assert results["failed"] == 0, results["errors"]
        assert results["updated"] == 3
        assert _labels(Folder.objects.get(name="Beta")) == {"Finance", "IT"}

    def test_malformed_internal_id_fails_only_its_row(
        self, api_client, domains, all_accessible
    ):
        wb = load_workbook(io.BytesIO(_export(api_client)))
        ws = wb.worksheets[0]
        for row in ws.iter_rows(min_row=2):
            if row[1].value == "Alpha":
                row[0].value = "not-a-uuid"
        ws.append([None, "Delta", None, None, None])
        edited = io.BytesIO()
        wb.save(edited)

        resp = _import(api_client, edited.getvalue(), "skip")
        assert resp.status_code == 200, resp.json()
        results = resp.json()["results"]
        assert results["failed"] == 1
        assert "Invalid internal_id" in str(results["errors"])
        assert results["created"] == 1
        assert Folder.objects.filter(name="Delta").exists()


def _iam_groups(folder) -> int:
    return UserGroup.objects.filter(folder=folder, builtin=True).count()


@pytest.mark.django_db
def test_iam_groups_round_trip(knox_admin_client):
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "create_iam_groups"])
    ws.append(["Grouped", "yes"])
    ws.append(["Plain", None])
    content = io.BytesIO()
    wb.save(content)

    resp = _import(knox_admin_client, content.getvalue(), "stop")
    assert resp.status_code == 200, resp.json()
    assert resp.json()["results"]["created"] == 2, resp.json()
    grouped = Folder.objects.get(name="Grouped")
    plain = Folder.objects.get(name="Plain")
    # Created by the import, the domain gets its groups as through the API.
    assert grouped.create_iam_groups and _iam_groups(grouped) > 0
    assert not plain.create_iam_groups and _iam_groups(plain) == 0

    wb = load_workbook(io.BytesIO(_export(knox_admin_client)))
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[header.index("name")].value == "Grouped":
            row[header.index("create_iam_groups")].value = False
    edited = io.BytesIO()
    wb.save(edited)

    resp = _import(knox_admin_client, edited.getvalue(), "update")
    assert resp.status_code == 200, resp.json()
    assert resp.json()["results"]["failed"] == 0, resp.json()
    grouped.refresh_from_db()
    assert not grouped.create_iam_groups
    assert _iam_groups(grouped) == 0
