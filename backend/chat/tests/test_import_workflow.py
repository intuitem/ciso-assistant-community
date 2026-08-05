"""Tests for tabular digestion, the document import workflow, and the
consumer-backed importer (dry-run + apply)."""

import io
import uuid

import pytest
from django.contrib.contenttypes.models import ContentType
from openpyxl import Workbook

from chat.tabular import (
    TABULAR_CONTENT_TYPES,
    detect_target,
    digest_document,
    extract_records,
    header_signature,
    map_columns,
)
from chat.workflows.base import WorkflowContext
from chat.workflows.import_document import ImportDocumentWorkflow

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CONTROL_ROWS = [
    ["Name", "Description", "Status", "Impact", "Ref_Id", "Custom Col"],
    ["MFA", "Enforce MFA", "active", 3, "CTL-1", "x"],
    ["Backups", "Daily backups", "in_progress", 4, "CTL-2", "y"],
]

RISK_HEADERS = [
    "ref_id",
    "name",
    "description",
    "current_probability",
    "current_impact",
    "treatment",
    "assets",
    "applied_controls",
]
RISK_ROWS = [
    RISK_HEADERS,
    [
        "R-1",
        "Ransomware",
        "Production data encrypted",
        "High",
        "High",
        "mitigate",
        "ERP|Backup server",
        "EDR rollout",
    ],
    ["R-2", "Data leak", "Exfiltration via SaaS", "Low", "Medium", "accept", "ERP", ""],
]

MATRIX_DEFINITION = {
    "probability": [
        {"id": 0, "name": "Low"},
        {"id": 1, "name": "Medium"},
        {"id": 2, "name": "High"},
    ],
    "impact": [
        {"id": 0, "name": "Low"},
        {"id": 1, "name": "Medium"},
        {"id": 2, "name": "High"},
    ],
    "risk": [
        {"id": 0, "name": "Low", "abbreviation": "L", "hexcolor": "#BBF7D0"},
        {"id": 1, "name": "Medium", "abbreviation": "M", "hexcolor": "#FDE047"},
        {"id": 2, "name": "High", "abbreviation": "H", "hexcolor": "#F87171"},
    ],
    "grid": [[0, 0, 1], [0, 1, 2], [1, 2, 2]],
}


class _StubFieldFile:
    def __init__(self, raw: bytes):
        self._raw = raw

    def open(self, mode="rb"):
        return io.BytesIO(self._raw)


class _StubDocument:
    def __init__(self, raw: bytes, content_type: str, filename: str):
        self.id = uuid.uuid4()
        self.file = _StubFieldFile(raw)
        self.content_type = content_type
        self.filename = filename


class _StubSession:
    def __init__(self):
        self.workflow_state = {}

    def save(self, update_fields=None):
        pass


def _xlsx_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _controls_doc() -> _StubDocument:
    return _StubDocument(_xlsx_bytes(CONTROL_ROWS), XLSX_CT, "controls.xlsx")


def _ctx(documents=None, message="", session=None, request=None, user=None):
    from chat.scoping import ReadScope

    return WorkflowContext(
        user_message=message,
        parsed_context=None,
        scope=ReadScope(user or (request.user if request else None)),
        llm=None,
        session=session,
        documents=documents or [],
        request=request,
    )


def _run(workflow, ctx):
    return list(workflow.run(ctx))


# ── DB fixtures ──────────────────────────────────────────────────────


@pytest.fixture
def admin_request(db):
    from core.apps import startup

    startup(sender=None)

    from rest_framework.test import APIRequestFactory

    from iam.models import User, UserGroup

    admin = User.objects.create_superuser("admin@chat-import.test")
    UserGroup.objects.get(name="BI-UG-ADM").user_set.add(admin)
    request = APIRequestFactory().post("/")
    request.user = admin
    return request


@pytest.fixture
def domain(admin_request):
    from iam.models import Folder

    return Folder.objects.create(
        name="Chat Import Tests",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )


@pytest.fixture
def controls_document(domain):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from chat.models import IndexedDocument

    return IndexedDocument.objects.create(
        folder=domain,
        file=SimpleUploadedFile("controls.xlsx", _xlsx_bytes(CONTROL_ROWS)),
        filename="controls.xlsx",
        content_type=XLSX_CT,
        source_type=IndexedDocument.SourceType.CHAT,
    )


@pytest.fixture
def reader_request(domain):
    """A user who may read the domain but not write anything into it."""
    from rest_framework.test import APIRequestFactory

    from iam.models import Role, RoleAssignment, User, UserGroup

    user = User.objects.create_user("reader@chat-import.test", is_published=True)
    group = UserGroup.objects.create(name="readers", folder=domain)
    group.user_set.add(user)
    assignment = RoleAssignment.objects.create(
        user_group=group,
        role=Role.objects.get(name="BI-RL-AUD"),
        folder=domain,
        is_recursive=True,
    )
    assignment.perimeter_folders.add(domain)
    request = APIRequestFactory().post("/")
    request.user = user
    return request


@pytest.fixture
def risk_matrix(domain):
    from core.models import RiskMatrix

    return RiskMatrix.objects.create(
        name="3x3",
        folder=domain,
        json_definition=MATRIX_DEFINITION,
    )


@pytest.fixture
def risk_document(domain):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from chat.models import IndexedDocument

    return IndexedDocument.objects.create(
        folder=domain,
        file=SimpleUploadedFile("Q1 risks.xlsx", _xlsx_bytes(RISK_ROWS)),
        filename="Q1 risks.xlsx",
        content_type=XLSX_CT,
        source_type=IndexedDocument.SourceType.CHAT,
    )


@pytest.fixture
def chat_session(admin_request, domain):
    from chat.models import ChatSession

    return ChatSession.objects.create(owner=admin_request.user)


# ── tabular (pure) ───────────────────────────────────────────────────


def test_digest_xlsx():
    digest = digest_document(_controls_doc())
    assert not digest.error
    assert digest.headers[0] == "Name"
    assert digest.normalized_headers == [
        "name",
        "description",
        "status",
        "impact",
        "ref_id",
        "custom col",
    ]
    assert digest.row_count == 2
    assert digest.sample_rows[0]["name"] == "MFA"
    assert digest.signature == header_signature(digest.normalized_headers)


def test_digest_csv_semicolon():
    raw = b"name;severity;status\nSQLi;critical;identified\n"
    doc = _StubDocument(raw, "text/csv", "findings.csv")
    digest = digest_document(doc)
    assert not digest.error
    assert digest.normalized_headers == ["name", "severity", "status"]
    assert digest.row_count == 1


def test_digest_empty_file_errors():
    doc = _StubDocument(_xlsx_bytes([]), XLSX_CT, "empty.xlsx")
    assert digest_document(doc).error


def test_signature_order_insensitive():
    assert header_signature(["b", "a"]) == header_signature(["a", "b"])
    assert header_signature(["a"]) != header_signature(["a", "b"])


def test_map_columns_aliases():
    mapped, unmapped = map_columns(["name", "impact", "custom col"], "applied_control")
    assert mapped == {"name": "name", "impact": "control_impact"}
    assert unmapped == ["custom col"]


def test_extract_records_keeps_source_headers_and_drops_unknown():
    mapped, _ = map_columns(
        ["name", "description", "status", "impact", "ref_id", "custom col"],
        "applied_control",
    )
    records = extract_records(_controls_doc(), mapped)
    assert len(records) == 2
    assert records[0]["name"] == "MFA"
    # Source header is preserved — the consumer resolves the alias itself.
    assert records[0]["impact"] == 3
    assert "custom col" not in records[0]


def test_extract_records_keeps_every_cost_column():
    headers = [
        "name",
        "cost_currency",
        "cost_build_fixed",
        "cost_run_fixed",
    ]
    doc = _StubDocument(
        _xlsx_bytes([headers, ["MFA", "EUR", 1000, 200]]), XLSX_CT, "cost.xlsx"
    )
    mapped, unmapped = map_columns(headers, "applied_control")
    assert not unmapped
    record = extract_records(doc, mapped)[0]
    # All three must survive: they share the canonical key "cost", so renaming
    # would collapse them into one value and drop the cost data entirely.
    assert record["cost_currency"] == "EUR"
    assert record["cost_build_fixed"] == 1000
    assert record["cost_run_fixed"] == 200


def test_detect_target_prefers_applied_control():
    headers = ["name", "status", "impact", "effort", "ref_id", "eta"]
    scores = detect_target(headers)
    assert scores[0][0] == "applied_control"


def test_detect_target_prefers_asset():
    headers = ["name", "type", "parent_assets", "ref_id"]
    scores = detect_target(headers)
    assert scores[0][0] == "asset"


def test_detect_target_prefers_risk_scenario():
    scores = detect_target(RISK_HEADERS)
    assert scores[0][0] == "risk_scenario"


def test_detect_target_still_prefers_finding():
    headers = ["ref_id", "name", "severity", "status", "eta", "observation"]
    scores = detect_target(headers)
    assert scores[0][0] == "finding"


def test_risk_columns_map_export_and_legacy_names():
    mapped, _ = map_columns(
        ["current_probability", "residual_proba", "additional_controls", "asset"],
        "risk_scenario",
    )
    assert mapped["current_probability"] == "current_proba"
    assert mapped["residual_proba"] == "residual_proba"
    assert mapped["additional_controls"] == "applied_controls"
    assert mapped["asset"] == "assets"


def test_container_fk_and_computed_levels_are_never_mapped():
    mapped, unmapped = map_columns(
        ["risk_assessment", "current_level", "name"], "risk_scenario"
    )
    assert mapped == {"name": "name"}
    assert set(unmapped) == {"risk_assessment", "current_level"}


def test_permission_model_is_the_container_when_there_is_one():
    from chat.tabular import permission_model_name

    assert permission_model_name("risk_scenario") == "riskassessment"
    assert permission_model_name("finding") == "findingsassessment"
    assert permission_model_name("asset") == "asset"


# ── workflow turns without a document (pure) ─────────────────────────


def test_no_document_no_state_guides_user():
    events = _run(ImportDocumentWorkflow(), _ctx(message="import my file"))
    assert any(e.type == "token" for e in events)


def test_ambiguous_headers_ask_for_target():
    raw = _xlsx_bytes([["name", "status", "description"], ["a", "b", "c"]])
    doc = _StubDocument(raw, XLSX_CT, "ambiguous.xlsx")
    session = _StubSession()
    events = _run(ImportDocumentWorkflow(), _ctx(documents=[doc], session=session))
    assert any(e.type == "pending_choice" for e in events)
    assert session.workflow_state["step"] == "awaiting_target"


def test_cancel_clears_state():
    session = _StubSession()
    session.workflow_state = {
        "workflow": "import_document",
        "step": "import_review",
        "data": {},
    }
    events = _run(ImportDocumentWorkflow(), _ctx(message="cancel", session=session))
    assert session.workflow_state == {}
    assert any(e.type == "token" for e in events)


def test_tabular_content_types_match_upload_validation():
    from chat.upload_validation import CHAT_UPLOAD_CONTENT_TYPES

    assert TABULAR_CONTENT_TYPES <= set(CHAT_UPLOAD_CONTENT_TYPES.values())


# ── importer (DB) ────────────────────────────────────────────────────


@pytest.mark.django_db
class TestImporter:
    def _mapping(self):
        mapped, _ = map_columns(
            ["name", "description", "status", "impact", "ref_id", "custom col"],
            "applied_control",
        )
        return mapped

    def test_dry_run_counts_without_writing(
        self, admin_request, domain, controls_document
    ):
        from chat.importer import run_import
        from core.models import AppliedControl

        report = run_import(
            admin_request,
            controls_document,
            self._mapping(),
            "applied_control",
            folder_id=str(domain.id),
            dry_run=True,
        )
        assert report["created"] == 2
        assert report["updated"] == 0
        assert report["failed"] == 0
        assert AppliedControl.objects.count() == 0

    def test_apply_then_reapply_updates(self, admin_request, domain, controls_document):
        from chat.importer import run_import
        from core.models import AppliedControl

        report = run_import(
            admin_request,
            controls_document,
            self._mapping(),
            "applied_control",
            folder_id=str(domain.id),
            dry_run=False,
        )
        assert report["created"] == 2
        controls = AppliedControl.objects.filter(folder=domain)
        assert controls.count() == 2
        assert controls.get(ref_id="CTL-1").name == "MFA"

        # Re-importing the same file matches on ref_id and updates in place.
        report = run_import(
            admin_request,
            controls_document,
            self._mapping(),
            "applied_control",
            folder_id=str(domain.id),
            dry_run=False,
        )
        assert report["created"] == 0
        assert report["updated"] == 2
        assert AppliedControl.objects.filter(folder=domain).count() == 2

    def test_findings_update_existing_assessment(self, admin_request, domain):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from chat.importer import run_import
        from chat.models import IndexedDocument
        from core.models import Finding, FindingsAssessment

        assessment = FindingsAssessment.objects.create(name="Q1 Pentest", folder=domain)
        Finding.objects.create(
            findings_assessment=assessment,
            name="Old SQLi",
            ref_id="F-1",
            folder=domain,
        )

        doc = IndexedDocument.objects.create(
            folder=domain,
            file=SimpleUploadedFile(
                "findings.xlsx",
                _xlsx_bytes(
                    [
                        ["ref_id", "name", "severity"],
                        ["F-1", "SQL injection (retested)", "high"],
                        ["F-2", "XSS on login", "medium"],
                    ]
                ),
            ),
            filename="findings.xlsx",
            content_type=XLSX_CT,
            source_type=IndexedDocument.SourceType.CHAT,
        )
        mapped, _ = map_columns(["ref_id", "name", "severity"], "finding")

        report = run_import(
            admin_request,
            doc,
            mapped,
            "finding",
            folder_id=str(domain.id),
            dry_run=False,
            target_id=str(assessment.id),
        )
        assert report["updated"] == 1
        assert report["created"] == 1
        # No new assessment was created; both findings live in the target.
        assert FindingsAssessment.objects.count() == 1
        assert assessment.findings.count() == 2
        assert assessment.findings.get(ref_id="F-1").name == "SQL injection (retested)"

    def test_findings_capture_the_asset_named_on_the_row(self, admin_request, domain):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from chat.importer import run_import
        from chat.models import IndexedDocument
        from core.models import Asset, FindingsAssessment

        Asset.objects.create(name="ERP", folder=domain)
        headers = ["ref_id", "name", "severity", "asset"]
        doc = IndexedDocument.objects.create(
            folder=domain,
            file=SimpleUploadedFile(
                "findings.xlsx",
                _xlsx_bytes(
                    [
                        headers,
                        ["F-1", "SQL injection", "high", "ERP"],
                        ["F-2", "Weak TLS", "medium", "Payment gateway|Load balancer"],
                    ]
                ),
            ),
            filename="findings.xlsx",
            content_type=XLSX_CT,
            source_type=IndexedDocument.SourceType.CHAT,
        )
        mapped, _ = map_columns(headers, "finding")
        assert mapped["asset"] == "asset"

        report = run_import(
            admin_request,
            doc,
            mapped,
            "finding",
            folder_id=str(domain.id),
            dry_run=False,
        )
        assert report["created"] == 2
        assert report["details"]["assets_created"] == 1

        assessment = FindingsAssessment.objects.get()
        assert assessment.findings.get(ref_id="F-1").asset.name == "ERP"
        gateway = assessment.findings.get(ref_id="F-2").asset
        assert gateway.name == "Payment gateway"
        assert gateway.folder == domain
        assert gateway.type == Asset.Type.SUPPORT

        # Finding.asset is a single FK: the second name leaves no orphan.
        assert not Asset.objects.filter(name="Load balancer").exists()
        assert Asset.objects.count() == 2

    def test_risk_scenarios_create_assessment_assets_and_controls(
        self, admin_request, domain, risk_matrix, risk_document
    ):
        from chat.importer import run_import
        from core.models import AppliedControl, Asset, RiskAssessment, RiskScenario

        mapped, _ = map_columns(RISK_HEADERS, "risk_scenario")
        report = run_import(
            admin_request,
            risk_document,
            mapped,
            "risk_scenario",
            folder_id=str(domain.id),
            dry_run=False,
            matrix_id=str(risk_matrix.id),
        )
        assert report["created"] == 2
        assert report["failed"] == 0
        assert report["details"]["assets_created"] == 2
        assert report["details"]["applied_controls_created"] == 1

        assessment = RiskAssessment.objects.get()
        # Named after the uploaded file, not a timestamp.
        assert assessment.name == "Q1 risks"
        assert assessment.risk_matrix == risk_matrix

        ransomware = RiskScenario.objects.get(ref_id="R-1")
        assert ransomware.current_proba == 2
        assert ransomware.current_impact == 2
        assert ransomware.treatment == "mitigate"
        assert set(ransomware.assets.values_list("name", flat=True)) == {
            "ERP",
            "Backup server",
        }
        assert list(ransomware.applied_controls.values_list("name", flat=True)) == [
            "EDR rollout"
        ]
        assert Asset.objects.filter(folder=domain).count() == 2
        assert AppliedControl.objects.filter(folder=domain).count() == 1

    def test_risk_scenarios_dry_run_writes_nothing(
        self, admin_request, domain, risk_matrix, risk_document
    ):
        from chat.importer import run_import
        from core.models import AppliedControl, Asset, RiskAssessment

        mapped, _ = map_columns(RISK_HEADERS, "risk_scenario")
        report = run_import(
            admin_request,
            risk_document,
            mapped,
            "risk_scenario",
            folder_id=str(domain.id),
            dry_run=True,
            matrix_id=str(risk_matrix.id),
        )
        assert report["created"] == 2
        assert RiskAssessment.objects.count() == 0
        assert Asset.objects.count() == 0
        assert AppliedControl.objects.count() == 0

    def test_risk_scenarios_update_existing_assessment(
        self, admin_request, domain, risk_matrix, risk_document
    ):
        from chat.importer import run_import
        from core.models import RiskAssessment, RiskScenario

        assessment = RiskAssessment.objects.create(
            name="Existing study", folder=domain, risk_matrix=risk_matrix
        )
        RiskScenario.objects.create(
            risk_assessment=assessment,
            name="Old ransomware",
            ref_id="R-1",
            folder=domain,
        )

        mapped, _ = map_columns(RISK_HEADERS, "risk_scenario")
        report = run_import(
            admin_request,
            risk_document,
            mapped,
            "risk_scenario",
            folder_id=str(domain.id),
            dry_run=False,
            target_id=str(assessment.id),
        )
        assert report["updated"] == 1
        assert report["created"] == 1
        assert RiskAssessment.objects.count() == 1
        assert assessment.risk_scenarios.count() == 2
        assert assessment.risk_scenarios.get(ref_id="R-1").name == "Ransomware"

    def test_reader_cannot_import_anything(
        self, reader_request, domain, risk_matrix, risk_document, controls_document
    ):
        from rest_framework.exceptions import PermissionDenied

        from chat.importer import check_import_permission, run_import
        from core.models import RiskAssessment

        assert not check_import_permission(
            reader_request.user, "risk_scenario", str(domain.id)
        )
        assert not check_import_permission(
            reader_request.user, "applied_control", str(domain.id)
        )

        mapped, _ = map_columns(RISK_HEADERS, "risk_scenario")
        with pytest.raises(PermissionDenied):
            run_import(
                reader_request,
                risk_document,
                mapped,
                "risk_scenario",
                folder_id=str(domain.id),
                dry_run=True,
                matrix_id=str(risk_matrix.id),
            )
        assert RiskAssessment.objects.count() == 0

    def test_assets_are_not_created_for_a_user_who_may_not_add_them(
        self, reader_request, domain
    ):
        from data_wizard.views import _resolve_assets
        from core.models import Asset

        resolved = _resolve_assets("Unknown ERP", domain, reader_request)

        assert resolved.ids == []
        assert resolved.created == []
        assert resolved.failed == ["Unknown ERP"]
        assert Asset.objects.count() == 0

    def test_missing_matrix_fails_the_whole_run(
        self, admin_request, domain, risk_document
    ):
        from chat.importer import run_import
        from core.models import RiskAssessment

        mapped, _ = map_columns(RISK_HEADERS, "risk_scenario")
        report = run_import(
            admin_request,
            risk_document,
            mapped,
            "risk_scenario",
            folder_id=str(domain.id),
            dry_run=True,
        )
        assert report["created"] == 0
        assert report["failed"] == 2
        assert RiskAssessment.objects.count() == 0


# ── workflow turns with a real document (DB) ─────────────────────────


@pytest.mark.django_db
class TestWorkflowWithDocument:
    def test_digest_turn_dry_runs_and_proposes(
        self, admin_request, domain, controls_document, chat_session
    ):
        from core.models import AppliedControl

        ctx = _ctx(
            documents=[controls_document],
            session=chat_session,
            request=admin_request,
        )
        events = _run(ImportDocumentWorkflow(), ctx)

        mapping_requests = [e for e in events if e.type == "mapping_request"]
        assert len(mapping_requests) == 1
        assert mapping_requests[0].content["signature"]

        narration = "".join(
            e.content
            for e in events
            if e.type == "token" and isinstance(e.content, str)
        )
        assert "controls.xlsx" in narration
        assert "custom col" in narration  # exceptions-first: ignored column named

        actions = [e for e in events if e.type == "pending_action"]
        assert len(actions) == 1
        assert actions[0].content["action"] == "import"
        assert actions[0].content["created"] == 2
        assert actions[0].content["updated"] == 0
        assert actions[0].content["folder_id"] == str(domain.id)

        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "import_review"
        state_data = chat_session.workflow_state["data"]
        assert state_data["target"] == "applied_control"
        assert state_data["mapping"]["impact"] == "control_impact"

        # Dry-run must not have written anything
        assert AppliedControl.objects.count() == 0

    def test_stated_target_beats_detection(
        self, admin_request, domain, controls_document, chat_session
    ):
        ctx = _ctx(
            documents=[controls_document],
            message="import these as pentest findings",
            session=chat_session,
            request=admin_request,
        )
        _run(ImportDocumentWorkflow(), ctx)
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["data"]["target"] == "finding"

    def test_risk_flow_asks_for_matrix_when_creating_an_assessment(
        self, admin_request, domain, risk_matrix, risk_document, chat_session
    ):
        from core.models import RiskAssessment, RiskMatrix

        RiskMatrix.objects.create(
            name="5x5", folder=domain, json_definition=MATRIX_DEFINITION
        )
        RiskAssessment.objects.create(
            name="Existing study", folder=domain, risk_matrix=risk_matrix
        )

        workflow = ImportDocumentWorkflow()
        events = _run(
            workflow,
            _ctx(
                documents=[risk_document],
                session=chat_session,
                request=admin_request,
            ),
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["data"]["target"] == "risk_scenario"
        assert chat_session.workflow_state["step"] == "awaiting_container"
        names = [i["name"] for i in events[-1].content["items"]]
        assert names == ["Create a new risk assessment", "Existing study"]

        events = _run(
            workflow, _ctx(message="new", session=chat_session, request=admin_request)
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "awaiting_matrix"
        assert {i["name"] for i in events[-1].content["items"]} == {"3x3", "5x5"}
        assert RiskAssessment.objects.count() == 1

        events = _run(
            workflow, _ctx(message="3x3", session=chat_session, request=admin_request)
        )
        chat_session.refresh_from_db()
        state_data = chat_session.workflow_state["data"]
        assert chat_session.workflow_state["step"] == "import_review"
        assert state_data["matrix_id"] == str(risk_matrix.id)

        action = [e for e in events if e.type == "pending_action"][-1]
        assert action.content["created"] == 2
        narration = "".join(
            e.content
            for e in events
            if e.type == "token" and isinstance(e.content, str)
        )
        assert "2 assets" in narration
        assert "1 applied controls" in narration
        assert RiskAssessment.objects.count() == 1

    def test_risk_flow_skips_the_matrix_question_for_an_existing_assessment(
        self, admin_request, domain, risk_matrix, risk_document, chat_session
    ):
        from core.models import RiskAssessment

        RiskAssessment.objects.create(
            name="Existing study", folder=domain, risk_matrix=risk_matrix
        )
        workflow = ImportDocumentWorkflow()
        _run(
            workflow,
            _ctx(
                documents=[risk_document], session=chat_session, request=admin_request
            ),
        )
        _run(
            workflow,
            _ctx(message="Existing study", session=chat_session, request=admin_request),
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "import_review"
        assert chat_session.workflow_state["data"]["container_name"] == "Existing study"

    def test_findings_flow_asks_for_container_then_updates(
        self, admin_request, domain, controls_document, chat_session
    ):
        from core.models import FindingsAssessment

        assessment = FindingsAssessment.objects.create(name="Q1 Pentest", folder=domain)

        workflow = ImportDocumentWorkflow()
        events = _run(
            workflow,
            _ctx(
                documents=[controls_document],
                message="import these as pentest findings",
                session=chat_session,
                request=admin_request,
            ),
        )
        choices = [e for e in events if e.type == "pending_choice"]
        assert len(choices) == 1
        names = [i["name"] for i in choices[0].content["items"]]
        assert "Q1 Pentest" in names
        assert "Create a new findings assessment" in names
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "awaiting_container"

        events = _run(
            workflow,
            _ctx(
                message="Q1 Pentest",
                session=chat_session,
                request=admin_request,
            ),
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "import_review"
        assert chat_session.workflow_state["data"]["container_id"] == str(assessment.id)
        actions = [e for e in events if e.type == "pending_action"]
        assert actions[0].content["target_name"] == "Q1 Pentest"

    def test_target_reply_resumes_and_dry_runs(
        self, admin_request, domain, chat_session
    ):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from chat.models import IndexedDocument

        doc = IndexedDocument.objects.create(
            folder=domain,
            file=SimpleUploadedFile(
                "ambiguous.xlsx",
                _xlsx_bytes([["name", "status", "description"], ["a", "b", "c"]]),
            ),
            filename="ambiguous.xlsx",
            content_type=XLSX_CT,
            source_type=IndexedDocument.SourceType.CHAT,
        )
        workflow = ImportDocumentWorkflow()
        _run(
            workflow,
            _ctx(
                documents=[doc],
                session=chat_session,
                request=admin_request,
            ),
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "awaiting_target"

        events = _run(
            workflow,
            _ctx(
                message="Assets",
                session=chat_session,
                request=admin_request,
            ),
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["step"] == "import_review"
        assert chat_session.workflow_state["data"]["target"] == "asset"
        assert any(e.type == "pending_action" for e in events)


# ── apply endpoint guards (DB) ───────────────────────────────────────


@pytest.mark.django_db
class TestApplyImportGuards:
    @pytest.fixture
    def api_client(self, admin_request):
        from knox.models import AuthToken
        from rest_framework.test import APIClient

        client = APIClient()
        client.credentials(
            HTTP_AUTHORIZATION=f"Token {AuthToken.objects.create(admin_request.user)[1]}"
        )
        return client

    def _stage(self, session, doc, **overrides):
        mapped, _ = map_columns(
            ["name", "description", "status", "impact", "ref_id", "custom col"],
            "applied_control",
        )
        session.workflow_state = {
            "workflow": "import_document",
            "step": "import_review",
            "data": {
                "document_id": str(doc.id),
                "target": "applied_control",
                "mapping": mapped,
                **overrides,
            },
        }
        session.save(update_fields=["workflow_state"])

    def _session_with_doc(self, chat_session, controls_document):
        controls_document.source_content_type = ContentType.objects.get_for_model(
            type(chat_session)
        )
        controls_document.source_object_id = chat_session.id
        controls_document.save()
        return chat_session

    def test_second_confirm_is_rejected(
        self, api_client, chat_session, controls_document, domain
    ):
        from core.models import AppliedControl

        session = self._session_with_doc(chat_session, controls_document)
        self._stage(session, controls_document)
        url = f"/api/chat/sessions/{session.id}/import/"

        first = api_client.post(url, {}, format="json")
        assert first.status_code == 200, first.data
        assert AppliedControl.objects.filter(folder=domain).count() == 2

        second = api_client.post(url, {}, format="json")
        assert second.status_code == 400
        assert AppliedControl.objects.filter(folder=domain).count() == 2

    def test_malformed_state_returns_400(
        self, api_client, chat_session, controls_document
    ):
        session = self._session_with_doc(chat_session, controls_document)
        self._stage(session, controls_document)
        state = session.workflow_state
        del state["data"]["mapping"]
        session.workflow_state = state
        session.save(update_fields=["workflow_state"])

        resp = api_client.post(
            f"/api/chat/sessions/{session.id}/import/", {}, format="json"
        )
        assert resp.status_code == 400

    def test_unknown_target_returns_400(
        self, api_client, chat_session, controls_document
    ):
        session = self._session_with_doc(chat_session, controls_document)
        self._stage(session, controls_document, target="no_such_model")

        resp = api_client.post(
            f"/api/chat/sessions/{session.id}/import/", {}, format="json"
        )
        assert resp.status_code == 400


# ── review fixes: resume gating, container match, folder pinning ─────


def test_should_resume_only_intercepts_answering_turns():
    from chat.workflows.import_document import should_resume

    staged = {"workflow": "import_document", "step": "import_review", "data": {}}
    asking = {"workflow": "import_document", "step": "awaiting_target", "data": {}}

    # An unrelated question must not re-enter the workflow (and re-run the dry-run).
    assert not should_resume(staged, "what is ISO 27001?")
    assert should_resume(staged, "cancel")
    assert should_resume(asking, "what is ISO 27001?")
    assert not should_resume(None, "cancel")
    assert not should_resume({"workflow": "ebios_rm_assist"}, "cancel")


def test_is_cancel_requires_whole_word():
    from chat.workflows.import_document import is_cancel

    assert is_cancel("cancel")
    assert is_cancel("Annuler s'il te plaît")
    assert not is_cancel("stopwatch inventory")


def test_import_review_turn_does_not_redo_dry_run():
    session = _StubSession()
    session.workflow_state = {
        "workflow": "import_document",
        "step": "import_review",
        "data": {"target": "applied_control"},
    }
    # No request/document available: a re-run would raise, this must not re-run.
    events = _run(ImportDocumentWorkflow(), _ctx(message="anything", session=session))
    assert not any(e.type == "pending_action" for e in events)
    assert session.workflow_state["step"] == "import_review"


@pytest.mark.django_db
class TestContainerMatching:
    def _stage_awaiting_container(self, session, candidates):
        session.workflow_state = {
            "workflow": "import_document",
            "step": "awaiting_container",
            "data": {
                "target": "finding",
                "document_id": "00000000-0000-0000-0000-000000000000",
                "normalized_headers": ["name"],
                "container_candidates": candidates,
            },
        }

    def test_candidate_name_containing_new_wins_over_keyword(
        self, admin_request, domain, chat_session, controls_document
    ):
        from core.models import FindingsAssessment

        assessment = FindingsAssessment.objects.create(name="Renewal Q1", folder=domain)
        workflow = ImportDocumentWorkflow()
        self._stage_awaiting_container(
            chat_session, [{"id": str(assessment.id), "name": "Renewal Q1"}]
        )
        chat_session.workflow_state["data"]["document_id"] = str(controls_document.id)
        chat_session.save(update_fields=["workflow_state"])

        # "Renewal" contains "new" — the assessment must still be selected.
        list(
            workflow.run(
                _ctx(
                    message="Renewal Q1",
                    session=chat_session,
                    request=admin_request,
                )
            )
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["data"]["container_id"] == str(assessment.id)

    def test_explicit_new_creates_container(
        self, admin_request, domain, chat_session, controls_document
    ):
        workflow = ImportDocumentWorkflow()
        self._stage_awaiting_container(
            chat_session, [{"id": "abc", "name": "Renewal Q1"}]
        )
        chat_session.workflow_state["data"]["document_id"] = str(controls_document.id)
        chat_session.save(update_fields=["workflow_state"])

        list(
            workflow.run(
                _ctx(
                    message="create a new one",
                    session=chat_session,
                    request=admin_request,
                )
            )
        )
        chat_session.refresh_from_db()
        assert chat_session.workflow_state["data"]["container_id"] is None


@pytest.mark.django_db
def test_dry_run_pins_folder_into_state(
    admin_request, domain, controls_document, chat_session
):
    events = _run(
        ImportDocumentWorkflow(),
        _ctx(
            documents=[controls_document],
            session=chat_session,
            request=admin_request,
        ),
    )
    chat_session.refresh_from_db()
    assert chat_session.workflow_state["data"]["folder_id"] == str(domain.id)
    card = next(e for e in events if e.type == "pending_action").content
    # No selector: switching folders after the dry-run would invalidate counts.
    assert card["available_folders"] == []
    assert card["folder_id"] == str(domain.id)
    assert card["truncated"] is False
